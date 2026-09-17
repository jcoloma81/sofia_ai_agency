#!/usr/bin/env python3
"""
Sofía Bridge — Conector de Escritorio para Microsoft Excel
=========================================================
Permite que Sofía escriba y actualice en tiempo real la planilla de Excel
que el comerciante tiene abierta en su computadora.

Características:
- Detecta si Excel está abierto en Windows (vía COM/win32com) y actualiza la pantalla en vivo.
- Soporta fallback universal con openpyxl (Linux, Mac y Windows con archivo guardado).
- Búsqueda inteligente de productos (coincidencia aproximada/fuzzy).
- Actualización de celdas individuales, agregado de filas de venta y lote de aumentos.
- Resaltado visual en celdas modificadas.
- Envío automático de confirmación (ACK) a Sofía en la nube.
"""

import os
import sys
import time
import json
import argparse
import logging
from datetime import datetime
from difflib import SequenceMatcher
from typing import Optional, Dict, Any, List, Tuple
import httpx
import openpyxl
from openpyxl.styles import PatternFill, Font

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%H:%M:%S"
)
logger = logging.getLogger("SofiaBridge")

# Colores suaves para resaltar cambios en Excel
YELLOW_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
GREEN_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

def normalize_text(text: str) -> str:
    """Normaliza texto quitando tildes y caracteres especiales."""
    t = str(text or "").lower().strip()
    replacements = (
        ("á", "a"), ("é", "e"), ("í", "i"), ("ó", "o"), ("ú", "u"),
        ("ü", "u"), ("ñ", "n")
    )
    for a, b in replacements:
        t = t.replace(a, b)
    return t

def similarity(a: str, b: str) -> float:
    """Calcula la similitud semántica simple entre dos nombres."""
    na, nb = normalize_text(a), normalize_text(b)
    if na in nb or nb in na:
        return 0.85
    words_a = set(na.split())
    words_b = set(nb.split())
    common = words_a.intersection(words_b)
    if common and (len(common) / max(len(words_a), 1)) >= 0.5:
        return 0.8
    return SequenceMatcher(None, na, nb).ratio()

def clean_alphanumeric(text: str) -> str:
    """Remueve puntuaciones, puntos, barras y espacios para comparar siglas (ej: 'p.c.' -> 'pc')."""
    norm = normalize_text(text)
    import re
    return re.sub(r'[^a-z0-9$]', '', norm)

def detect_column_role(header_text: str) -> Optional[str]:
    """
    Detecta el rol semántico de una columna a partir de su encabezado, soportando
    abreviaturas típicas de mostrador argentino (P.C., P.V., STK, ART., etc.).
    """
    raw = str(header_text or "").strip()
    if not raw:
        return None
    norm = normalize_text(raw)
    clean = clean_alphanumeric(raw)

    # 1. Total / Subtotal
    if clean in ("total", "subtotal", "importe", "monto") or "total" in norm:
        return "total"

    # 2. Fecha / Hora
    if clean in ("fecha", "date", "dia") or "fecha" in norm:
        return "fecha"
    if clean in ("hora", "time", "horario") or "hora" in norm:
        return "hora"

    # 3. Costo (DEBE CHEQUEARSE ANTES DE PRECIO porque 'precio costo' o 'p.c.' contienen la palabra 'precio')
    cost_exact_clean = {
        "pc", "cost", "costo", "pcompra", "pcosto", "prcosto", "preciocosto",
        "costosiva", "costociva", "costosiva", "costociva", "costociva", "costou", "costounitario"
    }
    cost_tokens = [
        "costo", "p.costo", "p. costo", "cost", "compra", "p.compra", "p. compra",
        "adquisicion", "fabrica", "coste", "proveedor", "mayorista"
    ]
    if clean in cost_exact_clean or any(t in norm for t in cost_tokens):
        return "cost"

    # 4. Precio de venta
    price_exact_clean = {
        "pv", "pvp", "pr", "precio", "precioventa", "pventa", "prventa", "publico",
        "ppub", "ppublico", "l1", "lista1", "lista", "valor", "$", "pvpiva", "pventasiva"
    }
    price_tokens = [
        "venta", "p.venta", "p. venta", "precio", "pvp", "p.v.p", "valor", "publico",
        "al publico", "mostrador", "lista 1", "lista", "contado", "p.contado", "$"
    ]
    if clean in price_exact_clean or any(t in norm for t in price_tokens):
        return "price"

    # 5. Stock / Existencias
    stock_exact_clean = {
        "stk", "stock", "cant", "cantidad", "disp", "disponible",
        "un", "unid", "unidades", "u", "saldo", "existencia", "existencias"
    }
    stock_tokens = [
        "stock", "cantidad", "cant", "disponible", "disp", "existencia",
        "existencias", "saldo", "unidades"
    ]
    if clean in stock_exact_clean or any(t in norm for t in stock_tokens):
        return "stock"

    # 6. Código / SKU / Barras
    code_exact_clean = {"cod", "codigo", "sku", "id", "ref", "ean", "cb"}
    code_tokens = ["codigo", "cod.", "sku", "referencia", "barra", "barras", "ean"]
    if clean in code_exact_clean or any(t in norm for t in code_tokens):
        return "code"

    # 7. Producto / Descripción / Detalle
    name_exact_clean = {
        "art", "articulo", "desc", "descripcion", "detalle", "nombre",
        "item", "producto", "prod", "concepto", "mercaderia", "material"
    }
    name_tokens = [
        "producto", "articulo", "descripcion", "detalle", "nombre",
        "concepto", "denominacion", "mercaderia", "material"
    ]
    if clean in name_exact_clean or any(t in norm for t in name_tokens):
        return "name"

    # 8. Rubro / Categoría / Marca
    cat_tokens = ["rubro", "categoria", "marca", "familia", "linea", "grupo"]
    if any(t in norm for t in cat_tokens):
        return "category"

    return None

def _add_convenience_aliases(mapping: Dict[str, int]) -> Dict[str, int]:
    """Copia roles canónicos a sinónimos frecuentes para búsquedas directas."""
    if "name" in mapping:
        for k in ["producto", "articulo", "descripcion", "detalle", "nombre", "item"]:
            mapping[k] = mapping["name"]
    if "cost" in mapping:
        for k in ["costo", "cost_price", "p.costo", "pc", "cost", "compra"]:
            mapping[k] = mapping["cost"]
    if "price" in mapping:
        for k in ["precio", "sale_price", "p.venta", "pv", "pvp", "valor", "$"]:
            mapping[k] = mapping["price"]
    if "stock" in mapping:
        for k in ["cantidad", "cant", "stk", "disp", "unidades"]:
            mapping[k] = mapping["stock"]
    if "code" in mapping:
        for k in ["codigo", "cod", "sku", "id"]:
            mapping[k] = mapping["code"]
    return mapping

class ExcelOperator:
    """Maneja la lectura y modificación de archivos Excel con tolerancia a cualquier formato."""

    def __init__(self, file_path: str, highlight: bool = True):
        self.file_path = os.path.abspath(file_path)
        self.highlight = highlight
        self.is_windows = sys.platform.startswith("win")
        self._ensure_file_exists()

    def _ensure_file_exists(self):
        """Si el archivo no existe, crea una plantilla de ejemplo."""
        if not os.path.exists(self.file_path):
            logger.info(f"Creando planilla inicial en {self.file_path}...")
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Precios"
            headers = ["Código", "Descripción", "Costo", "Precio Venta", "Stock"]
            ws.append(headers)
            header_font = Font(bold=True)
            for col in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col)
                cell.font = header_font
            
            # Datos de muestra tipo ferretería
            sample_data = [
                ["ART-01", "Cal Hidratada Loma Negra 25kg", 3200, 4500, 25],
                ["ART-02", "Cemento Loma Negra 50kg", 7500, 10500, 40],
                ["ART-03", "Curva PVC 110mm Tigre", 5500, 7800, 15],
                ["ART-04", "Pegamento PVC 125cc Tigre", 2200, 3100, 18],
                ["ART-05", "Lija al Agua Grano 100", 450, 650, 100],
                ["ART-06", "Tornillo Fix 4x40mm (Caja x 100)", 1800, 2600, 30],
                ["ART-07", "Cinta Aisladora Negra 20m", 950, 1400, 50],
            ]
            for row in sample_data:
                ws.append(row)

            # Hoja de Ventas
            ws_v = wb.create_sheet(title="Ventas")
            v_headers = ["Fecha", "Hora", "Detalle de Venta", "Total"]
            ws_v.append(v_headers)
            for col in range(1, len(v_headers) + 1):
                ws_v.cell(row=1, column=col).font = header_font

            wb.save(self.file_path)
            logger.info(f"✅ Planilla demo creada con éxito: {self.file_path}")

    def _try_win32_com(self) -> Tuple[bool, Any]:
        """Intenta conectarse a una instancia activa de Microsoft Excel en Windows."""
        if not self.is_windows:
            return False, None
        try:
            import win32com.client
            excel = win32com.client.GetObject(Class="Excel.Application")
            for wb in excel.Workbooks:
                if os.path.basename(wb.FullName).lower() == os.path.basename(self.file_path).lower():
                    return True, (excel, wb)
        except Exception:
            pass
        return False, None

    def _resolve_sheet_openpyxl(self, wb: openpyxl.Workbook, sheet_name: Optional[str] = None, purpose: str = "prices"):
        """Resuelve inteligentemente la hoja adecuada en openpyxl según su nombre o propósito."""
        if sheet_name and sheet_name in wb.sheetnames:
            return wb[sheet_name]

        if purpose == "sales":
            for sname in wb.sheetnames:
                ns = normalize_text(sname)
                if any(k in ns for k in ["venta", "movimiento", "caja", "diario", "operacion"]):
                    return wb[sname]
            # Si no existe, crear la hoja de ventas automáticamente
            ws_v = wb.create_sheet(title=sheet_name or "Ventas")
            v_headers = ["Fecha", "Hora", "Detalle de Venta", "Total"]
            ws_v.append(v_headers)
            header_font = Font(bold=True)
            for col in range(1, len(v_headers) + 1):
                ws_v.cell(row=1, column=col).font = header_font
            return ws_v

        # Para catálogo/precios
        if sheet_name:
            for sname in wb.sheetnames:
                if normalize_text(sheet_name) in normalize_text(sname) or normalize_text(sname) in normalize_text(sheet_name):
                    return wb[sname]

        for sname in wb.sheetnames:
            ns = normalize_text(sname)
            if any(k in ns for k in ["precio", "articulo", "producto", "stock", "catalogo", "lista", "inventario", "mercaderia"]):
                return wb[sname]

        return wb.active

    def _resolve_sheet_com(self, wb: Any, sheet_name: Optional[str] = None, purpose: str = "prices"):
        """Resuelve inteligentemente la hoja adecuada en win32com según su nombre o propósito."""
        if sheet_name:
            try:
                return wb.Sheets(sheet_name)
            except Exception:
                pass

        try:
            keywords = ["venta", "movimiento", "caja", "diario"] if purpose == "sales" else ["precio", "articulo", "producto", "stock", "catalogo", "lista"]
            for i in range(1, wb.Sheets.Count + 1):
                s = wb.Sheets(i)
                ns = normalize_text(s.Name)
                if any(k in ns for k in keywords):
                    return s
        except Exception:
            pass

        return wb.ActiveSheet

    def _find_header_row_and_mapping(self, sheet: Any) -> Tuple[int, Dict[str, int]]:
        """
        Escanea las primeras 10 filas de la hoja para detectar dónde arrancan los encabezados
        reales (saltando logos, títulos vacíos o fechas) y mapear las columnas.
        """
        best_row = 1
        best_mapping = {}
        best_score = 0
        max_scan = min(sheet.max_row, 10) if hasattr(sheet, "max_row") else 10

        for r in range(1, max_scan + 1):
            mapping = {}
            score = 0
            max_c = sheet.max_column if hasattr(sheet, "max_column") else 30
            for c in range(1, max_c + 1):
                val = sheet.cell(row=r, column=c).value
                if val is not None:
                    role = detect_column_role(str(val))
                    if role and role not in mapping:
                        mapping[role] = c
                        if role in ("name", "price"):
                            score += 3
                        elif role in ("cost", "stock", "code"):
                            score += 2
                        else:
                            score += 1
            if score > best_score:
                best_score = score
                best_row = r
                best_mapping = mapping

        if not best_mapping or best_score < 2:
            # Fallback fila 1
            best_row = 1
            max_c = sheet.max_column if hasattr(sheet, "max_column") else 10
            headers = [str(sheet.cell(row=1, column=c).value or "") for c in range(1, max_c + 1)]
            best_mapping = self._detect_headers_list(headers)

        _add_convenience_aliases(best_mapping)
        return best_row, best_mapping

    def _find_header_row_and_mapping_com(self, sheet: Any, max_cols: int) -> Tuple[int, Dict[str, int]]:
        """Detector de encabezados y mapeo para win32com escaneando las primeras 10 filas."""
        best_row = 1
        best_mapping = {}
        best_score = 0
        max_scan = min(sheet.UsedRange.Rows.Count, 10)

        for r in range(1, max_scan + 1):
            mapping = {}
            score = 0
            for c in range(1, max_cols + 1):
                val = sheet.Cells(r, c).Value
                if val is not None:
                    role = detect_column_role(str(val))
                    if role and role not in mapping:
                        mapping[role] = c
                        if role in ("name", "price"):
                            score += 3
                        elif role in ("cost", "stock", "code"):
                            score += 2
                        else:
                            score += 1
            if score > best_score:
                best_score = score
                best_row = r
                best_mapping = mapping

        if not best_mapping or best_score < 2:
            best_row = 1
            headers = [str(sheet.Cells(1, c).Value or "") for c in range(1, max_cols + 1)]
            best_mapping = self._detect_headers_list(headers)

        _add_convenience_aliases(best_mapping)
        return best_row, best_mapping

    def append_row(self, values: List[Any], sheet_name: Optional[str] = None) -> str:
        """Agrega un renglón al final de la planilla (ej: venta)."""
        has_com, com_objs = self._try_win32_com()
        if has_com:
            excel, wb = com_objs
            try:
                sheet = self._resolve_sheet_com(wb, sheet_name=sheet_name, purpose="sales")
                last_row = sheet.UsedRange.Rows.Count + 1
                for idx, val in enumerate(values, start=1):
                    sheet.Cells(last_row, idx).Value = val
                    if self.highlight:
                        sheet.Cells(last_row, idx).Interior.Color = 0xE2EFDA  # Soft green
                return f"Fila #{last_row} agregada en vivo en Excel abierto ({len(values)} columnas)."
            except Exception as e:
                logger.warning(f"Fallo COM en vivo: {e}. Usando openpyxl fallback...")

        # Fallback OpenPyXL
        wb = openpyxl.load_workbook(self.file_path)
        sheet = self._resolve_sheet_openpyxl(wb, sheet_name=sheet_name, purpose="sales")
        sheet.append(values)
        last_row = sheet.max_row
        if self.highlight:
            for col in range(1, len(values) + 1):
                sheet.cell(row=last_row, column=col).fill = GREEN_FILL
        wb.save(self.file_path)
        return f"Fila #{last_row} agregada con éxito en '{sheet.title}'."

    def update_product(self, search_term: str, updates: Dict[str, Any], sheet_name: Optional[str] = None) -> str:
        """Busca un producto existente en el Excel y modifica sus campos (precio, stock, etc.)."""
        has_com, com_objs = self._try_win32_com()
        if has_com:
            excel, wb = com_objs
            try:
                sheet = self._resolve_sheet_com(wb, sheet_name=sheet_name, purpose="prices")
                max_r = sheet.UsedRange.Rows.Count
                max_c = sheet.UsedRange.Columns.Count

                # 1. Identificar encabezados dinámicamente escaneando filas iniciales
                header_row, col_map = self._find_header_row_and_mapping_com(sheet, max_c)
                
                # 2. Buscar fila del producto a partir de header_row + 1
                best_row = None
                best_score = 0.0
                name_col = col_map.get("name", 2)

                for r in range(header_row + 1, max_r + 1):
                    cell_val = str(sheet.Cells(r, name_col).Value or "")
                    score = similarity(search_term, cell_val)
                    if score > best_score:
                        best_score = score
                        best_row = r

                if best_row and best_score >= 0.5:
                    prod_name = sheet.Cells(best_row, name_col).Value
                    changed_fields = []
                    for k, v in updates.items():
                        norm_k = normalize_text(k)
                        target_col = col_map.get(norm_k) or col_map.get(detect_column_role(k))
                        if target_col:
                            sheet.Cells(best_row, target_col).Value = v
                            if self.highlight:
                                sheet.Cells(best_row, target_col).Interior.Color = 0xFFF2CC
                            changed_fields.append(f"{k} -> {v}")

                    return f"Fila #{best_row} actualizada en vivo ('{prod_name}'): {', '.join(changed_fields)}"
            except Exception as e:
                logger.warning(f"Fallo COM en vivo: {e}. Usando openpyxl fallback...")

        # Fallback OpenPyXL
        wb = openpyxl.load_workbook(self.file_path)
        sheet = self._resolve_sheet_openpyxl(wb, sheet_name=sheet_name, purpose="prices")

        # Detectar columnas inteligentemente
        header_row, col_map = self._find_header_row_and_mapping(sheet)

        best_row = None
        best_score = 0.0
        name_col = col_map.get("name", 2)

        for r in range(header_row + 1, sheet.max_row + 1):
            cell_val = str(sheet.cell(row=r, column=name_col).value or "")
            score = similarity(search_term, cell_val)
            if score > best_score:
                best_score = score
                best_row = r

        if best_row and best_score >= 0.5:
            prod_name = sheet.cell(row=best_row, column=name_col).value
            changed_fields = []
            for k, v in updates.items():
                norm_k = normalize_text(k)
                target_col = col_map.get(norm_k) or col_map.get(detect_column_role(k))
                if target_col:
                    cell = sheet.cell(row=best_row, column=target_col)
                    cell.value = v
                    if self.highlight:
                        cell.fill = YELLOW_FILL
                    changed_fields.append(f"{k}: {v}")

            wb.save(self.file_path)
            return f"Fila #{best_row} actualizada ('{prod_name}'): {', '.join(changed_fields)}"

        return f"No se encontró un producto similar a '{search_term}' en la planilla."

    def batch_update(self, items: List[Dict[str, Any]], sheet_name: Optional[str] = None) -> str:
        """Actualiza una lista de productos en lote (aumentos de proveedores)."""
        wb = openpyxl.load_workbook(self.file_path)
        sheet = self._resolve_sheet_openpyxl(wb, sheet_name=sheet_name, purpose="prices")

        header_row, col_map = self._find_header_row_and_mapping(sheet)
        name_col = col_map.get("name", 2)
        price_col = col_map.get("price", col_map.get("precio", 4))
        cost_col = col_map.get("cost", col_map.get("costo", 3))

        updated_count = 0
        added_count = 0

        for item in items:
            it_name = str(item.get("name", "")).strip()
            if not it_name:
                continue

            # Buscar fila existente desde header_row + 1
            matched_row = None
            best_score = 0.0
            for r in range(header_row + 1, sheet.max_row + 1):
                cell_val = str(sheet.cell(row=r, column=name_col).value or "")
                score = similarity(it_name, cell_val)
                if score > best_score:
                    best_score = score
                    matched_row = r

            if matched_row and best_score >= 0.55:
                if "cost_price" in item and cost_col:
                    sheet.cell(row=matched_row, column=cost_col).value = item["cost_price"]
                    if self.highlight:
                        sheet.cell(row=matched_row, column=cost_col).fill = YELLOW_FILL
                if "sale_price" in item and price_col:
                    sheet.cell(row=matched_row, column=price_col).value = item["sale_price"]
                    if self.highlight:
                        sheet.cell(row=matched_row, column=price_col).fill = YELLOW_FILL
                updated_count += 1
            else:
                # Agregar nuevo producto mapeando dinámicamente sus columnas
                target_cols = max(sheet.max_column, max(col_map.values(), default=5))
                new_row = [None] * target_cols
                if "code" in col_map:
                    new_row[col_map["code"] - 1] = item.get("code") or f"ART-{sheet.max_row}"
                if "name" in col_map:
                    new_row[col_map["name"] - 1] = it_name
                if "cost" in col_map:
                    new_row[col_map["cost"] - 1] = item.get("cost_price", 0)
                if "price" in col_map:
                    new_row[col_map["price"] - 1] = item.get("sale_price", 0)
                if "stock" in col_map:
                    new_row[col_map["stock"] - 1] = item.get("stock", 0)

                # Si no había mapeo suficiente, fallback estándar
                if not any(new_row):
                    new_row = [
                        item.get("code") or f"ART-{sheet.max_row}",
                        it_name,
                        item.get("cost_price", 0),
                        item.get("sale_price", 0),
                        0
                    ]

                sheet.append(new_row)
                if self.highlight:
                    for col in range(1, len(new_row) + 1):
                        sheet.cell(row=sheet.max_row, column=col).fill = GREEN_FILL
                added_count += 1

        wb.save(self.file_path)
        return f"Lote procesado: {updated_count} productos actualizados, {added_count} nuevos agregados."

    def _detect_headers_list(self, headers: List[str]) -> Dict[str, int]:
        mapping = {}
        for idx, h in enumerate(headers, start=1):
            role = detect_column_role(h)
            if role and role not in mapping:
                mapping[role] = idx
        _add_convenience_aliases(mapping)
        return mapping

    def _detect_headers_com(self, sheet: Any, max_cols: int) -> Dict[str, int]:
        headers = [str(sheet.Cells(1, c).Value or "") for c in range(1, max_cols + 1)]
        return self._detect_headers_list(headers)


class SofiaBridgeClient:
    """Cliente de polling y ejecución remota para la PC del comerciante."""

    def __init__(self, server_url: str, merchant_phone: str, file_path: str, poll_interval: int = 2, gemini_api_key: Optional[str] = None):
        self.server_url = server_url.rstrip("/")
        self.merchant_phone = "".join(filter(str.isdigit, merchant_phone))
        self.operator = ExcelOperator(file_path)
        self.poll_interval = poll_interval
        self.gemini_api_key = gemini_api_key
        self.client_version = "1.1.0"
        self.running = True

    def run(self):
        logger.info("=" * 60)
        logger.info("🚀 SOFÍA BRIDGE — Conector Activo para Microsoft Excel")
        logger.info(f"📱 Comerciante ID: {self.merchant_phone}")
        logger.info(f"📊 Planilla conectada: {self.operator.file_path}")
        logger.info(f"🌐 Servidor Sofía: {self.server_url}")
        logger.info("=" * 60)
        logger.info("Esperando órdenes desde WhatsApp... (Presioná Ctrl+C para salir)")

        with httpx.Client(timeout=10.0) as client:
            # 1. Ping inicial al servidor central
            try:
                ping_url = f"{self.server_url}/api/v1/bridge/ping"
                ping_resp = client.post(ping_url, json={
                    "merchant_phone": self.merchant_phone,
                    "version": self.client_version,
                    "gemini_api_key": self.gemini_api_key
                })
                if ping_resp.status_code == 200:
                    pdata = ping_resp.json()
                    logger.info(f"🟢 Servidor Central: {pdata.get('message', 'Online')} (v{pdata.get('server_version', '1.1.0')})")
                    if pdata.get("update_available"):
                        logger.warning(f"⚠️ Actualización disponible en el servidor: v{pdata.get('server_version')}")
            except Exception as pe:
                logger.warning(f"Conexión inicial offline ({pe}). Continuando en modo local...")

            while self.running:
                try:
                    self._check_and_execute(client)
                except KeyboardInterrupt:
                    logger.info("Saliendo de Sofía Bridge...")
                    break
                except Exception as e:
                    logger.warning(f"Error en ciclo de sincronización: {e}")
                time.sleep(self.poll_interval)

    def _check_and_execute(self, client: httpx.Client):
        url = f"{self.server_url}/api/v1/bridge/commands"
        resp = client.get(url, params={"merchant_phone": self.merchant_phone})
        if resp.status_code != 200:
            return

        data = resp.json()
        commands = data.get("commands", [])
        for cmd in commands:
            cmd_id = cmd["command_id"]
            action = cmd["action"]
            payload = cmd.get("payload", {})
            sheet = cmd.get("sheet_name")

            logger.info(f"⚡ [NUEVA ORDEN RECIBIDA] Acción: {action} (ID: {cmd_id[:8]}...)")
            result_msg = ""
            status = "completed"

            try:
                if action == "append_row":
                    # Si es una venta o registro libre
                    values = payload.get("values", [])
                    if not values:
                        # Convertir dict a fila estándar
                        values = [
                            payload.get("fecha", datetime.now().strftime("%d/%m")),
                            payload.get("hora", datetime.now().strftime("%H:%M")),
                            payload.get("detalle", ""),
                            payload.get("total", 0)
                        ]
                    result_msg = self.operator.append_row(values, sheet_name=sheet or "Ventas")

                elif action == "update_product":
                    search = payload.get("search") or payload.get("item") or ""
                    updates = payload.get("updates") or payload.get("campos") or {}
                    if not updates and "precio" in payload:
                        updates["precio"] = payload["precio"]
                    if "stock" in payload:
                        updates["stock"] = payload["stock"]
                    result_msg = self.operator.update_product(search, updates, sheet_name=sheet or "Precios")

                elif action == "batch_update":
                    items = payload.get("items", [])
                    result_msg = self.operator.batch_update(items, sheet_name=sheet or "Precios")

                else:
                    status = "failed"
                    result_msg = f"Acción desconocida: {action}"

            except Exception as ex:
                status = "failed"
                result_msg = f"Error ejecutando en Excel: {ex}"
                logger.error(result_msg)

            logger.info(f"✨ Resultado: {result_msg}")

            # Enviar confirmación (ACK) a Sofía
            ack_url = f"{self.server_url}/api/v1/bridge/ack"
            client.post(ack_url, json={
                "command_id": cmd_id,
                "status": status,
                "result_message": result_msg,
                "notify_merchant": True
            })


def main():
    # Buscar config.json en dir actual o junto al script
    script_dir = os.path.dirname(os.path.abspath(__file__))
    config_paths = ["config.json", os.path.join(script_dir, "config.json")]
    cfg = {}
    for cp in config_paths:
        if os.path.exists(cp):
            try:
                with open(cp, "r", encoding="utf-8") as f:
                    cfg = json.load(f)
                break
            except Exception:
                pass

    parser = argparse.ArgumentParser(description="Sofía Bridge para Microsoft Excel")
    parser.add_argument("--merchant", default=cfg.get("merchant_phone", "5493434991122"), help="Teléfono E.164 del comercio")
    parser.add_argument("--server", default=cfg.get("server_url", "https://sofia-ai-agency.onrender.com"), help="URL del backend de Sofía")
    parser.add_argument("--file", default=cfg.get("excel_path", "ferreteria_demo.xlsx"), help="Ruta al archivo Excel")
    parser.add_argument("--interval", type=int, default=cfg.get("poll_interval", 2), help="Intervalo de sondeo en segundos")
    args = parser.parse_args()

    gemini_key = cfg.get("gemini_api_key") or os.environ.get("GEMINI_API_KEY")

    client = SofiaBridgeClient(
        server_url=args.server,
        merchant_phone=args.merchant,
        file_path=args.file,
        poll_interval=args.interval,
        gemini_api_key=gemini_key
    )
    client.run()

if __name__ == "__main__":
    main()
