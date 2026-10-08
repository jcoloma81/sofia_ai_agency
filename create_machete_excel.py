import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# Style definitions
header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Navy Blue
header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

section_fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid") # Soft Blue
section_font = Font(name="Segoe UI", size=11, bold=True, color="1E3A8A")

accent_fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid") # Green
accent_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

bold_font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
normal_font = Font(name="Segoe UI", size=10, color="1E293B")
italic_font = Font(name="Segoe UI", size=9, italic=True, color="475569")

thin_border = Border(
    left=Side(style='thin', color="CBD5E1"),
    right=Side(style='thin', color="CBD5E1"),
    top=Side(style='thin', color="CBD5E1"),
    bottom=Side(style='thin', color="CBD5E1")
)

def style_header(row, fill=header_fill, font=header_font):
    for cell in row:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

def auto_fit_columns(ws, max_len_cap=70):
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val = str(cell.value or '')
            lines = val.split('\n')
            for l in lines:
                if len(l) > max_len:
                    max_len = len(l)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), max_len_cap)

# ==============================================================================
# HOJA 1: MACHETE DE CALLE & OBJECIONES
# ==============================================================================
ws1 = wb.active
ws1.title = "Machete de Venta y Objeciones"
ws1.views.sheetView[0].showGridLines = True

ws1.append(["SOFÍA AI — MACHETE DE VENTA EN CALLE Y PREGUNTAS CLAVE"])
ws1.merge_cells("A1:D1")
ws1["A1"].font = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")
ws1["A1"].alignment = Alignment(vertical="center")
ws1.row_dimensions[1].height = 28

ws1.append(["Guía rápida para responder dudas técnicas, operativas y comerciales en 1 minuto frente al dueño o director médico."])
ws1.merge_cells("A2:D2")
ws1["A2"].font = italic_font
ws1.row_dimensions[2].height = 18
ws1.append([])

headers1 = ["Pregunta / Situación del Cliente", "¿Cómo funciona técnicamente?", "Qué responderle al Cliente (Pitch de Impacto)", "Regla de Oro / Tip de Venta"]
ws1.append(headers1)
style_header(ws1[4])
ws1.row_dimensions[4].height = 25

rows_sheet1 = [
    (
        "¿Cómo le doy el WhatsApp a un paciente en el mostrador si no tiene hashtag?",
        "1. QR en recepción: abre el chat con texto preescrito en 1 segundo.\n2. Reconocimiento por base de datos: si el teléfono ya es paciente, Sofía sabe quién es de inmediato.\n3. Si es alguien 100% nuevo que tipea el número a mano, Sofía le muestra un menú interactivo con botones.",
        "«En el mostrador ponemos un código QR en acrílico: el paciente apunta la cámara de su celular y ya le abre el WhatsApp listo para agendar. No tiene que guardar números ni tipear nada.»",
        "Destacar la comodidad: la gente odia agendar 13 dígitos a mano. El QR es estándar y profesional."
    ),
    (
        "¿Qué pasa si en mi clínica tenemos 2 o 3 secretarias?",
        "Sofía se integra al Grupo de WhatsApp interno de Recepción. A las 15:00 hs el aviso llega al grupo y cualquiera de las secretarias puede poner 'mantener a [Nombre]' o audios.",
        "«Sofía conecta a todo el equipo. A las 15 hs le avisa al grupo de secretarias los turnos en duda; la secretaria de la mañana o de la tarde pueden mandar audios y Sofía actualiza la misma agenda.»",
        "Resuelve el clásico problema de turnos rotativos donde una secretaria no sabe qué hizo la otra."
    ),
    (
        "¿Por qué me conviene pasarles mi base de datos de pacientes en Excel al inicio?",
        "La base inicial asocia cada teléfono a la clínica. Cuando María escribe, Sofía la saluda por su nombre sin pedir hashtags ni menús, y arranca con los recordatorios de la semana que viene.",
        "«Nos pasás tu lista de pacientes y turnos de la semana próxima, y desde el minuto uno Sofía saluda a cada paciente por su nombre y empieza a salvar los turnos que antes se perdían.»",
        "El efecto sorpresa: cuando el paciente escribe y la IA lo llama por su nombre, la percepción de valor se multiplica por diez."
    ),
    (
        "¿Tengo que cambiar mi número de WhatsApp?",
        "No. Sofía opera desde la Central Oficial Verificada de Meta. El paciente recibe los recordatorios con el nombre, logo y dirección de la clínica sin invadir el WhatsApp de la dueña.",
        "«No cambiás nada. Tu WhatsApp personal o comercial sigue igual. Sofía es un canal oficial de soporte y turnos que trabaja para vos 24/7.»",
        "Para clientes grandes que exigen usar su número histórico, Meta permite conectar su número propio (Plan Enterprise)."
    ),
    (
        "¿Qué pasa si un paciente toca [NO, REPROGRAMAR]?",
        "Se ejecutan 3 pasos automáticos:\n1. Se agradece cordialmente al paciente y se cancela el turno.\n2. Se alerta por WhatsApp a la secretaria.\n3. Se activa la lista de espera (2 horas de margen) o el Adelanta-Turnos si la lista está vacía.",
        "«El sistema no se queda quieto: le agradece al paciente, le avisa a tu secretaria y sale a buscar automáticamente a un suplente en lista de espera o le adelanta el turno a quien esperaba para la semana que viene.»",
        "Un hueco no cubierto cuesta entre $25.000 y $50.000. Sofía lo cubre en minutos."
    ),
    (
        "¿Qué pasa si aumentan mis precios o cuotas?",
        "Modo Director por voz: el dueño o secretaria manda un audio de 5 segundos ('Sofi, la cuota subió a $45.000'). Sofía actualiza la base y cotiza con el nuevo valor.",
        "«Le mandás un audio de 5 segundos a Sofía como si fuera tu asistente personal y listo, queda actualizado en el momento.»",
        "Cero software complejo ni paneles difíciles. Todo se gestiona por WhatsApp."
    )
]

for r in rows_sheet1:
    ws1.append(list(r))
    curr_row = ws1.max_row
    ws1.row_dimensions[curr_row].height = 45
    for c_idx, cell in enumerate(ws1[curr_row], 1):
        cell.font = normal_font
        cell.border = thin_border
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if c_idx == 1:
            cell.font = bold_font

auto_fit_columns(ws1)

# ==============================================================================
# HOJA 2: KIT DE 5 PLANTILLAS META
# ==============================================================================
ws2 = wb.create_sheet(title="Kit 5 Plantillas Meta")
ws2.views.sheetView[0].showGridLines = True

ws2.append(["KIT MAESTRO DE PLANTILLAS UNIVERSALES PARA META WHATSAPP CLOUD API"])
ws2.merge_cells("A1:F1")
ws2["A1"].font = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")
ws2.row_dimensions[1].height = 28

ws2.append(["Con estas 5 plantillas maestras reutilizables cubrimos el 100% de los clientes y rubros. Se aprueban en minutos y tienen costo mínimo."])
ws2.merge_cells("A2:F2")
ws2["A2"].font = italic_font
ws2.row_dimensions[2].height = 18
ws2.append([])

headers2 = ["Nombre de Plantilla", "Categoría Meta", "Rubros que Aplica", "Texto del Mensaje (Variables dinámicas)", "Botones Interactivos", "Beneficio & Costo"]
ws2.append(headers2)
style_header(ws2[4])
ws2.row_dimensions[4].height = 25

rows_sheet2 = [
    (
        "recordatorio_turno_v2",
        "UTILITY (Utilidad)",
        "Salud, Odontología, Kinesiología, Estética, Veterinarias, Peluquerías",
        "«Hola {{1}}, te recordamos tu turno en {{2}} para el {{3}} a las {{4}} hs con {{5}}. Por favor confirmá tu asistencia antes de las 18:00 hs.»",
        "• [✅ SÍ, CONFIRMO]\n• [🔄 NO, REPROGRAMAR]",
        "Aprobación en 5 min. Tarifa Utility (~USD $0,015). Erradica el 25% de ausentismo."
    ),
    (
        "adelanta_turno_v1",
        "UTILITY (Utilidad)",
        "Salud y Servicios con Agenda cuando se libera un turno de mañana",
        "«Hola {{1}}, en {{2}} se liberó un turno para mañana {{3}} a las {{4}} hs con {{5}}. Como tu turno es para el {{6}}, ¿te gustaría adelantarlo para mañana?»",
        "• [⚡ SÍ, ADELANTAR]\n• [📅 NO, MANTENER]",
        "Llena huecos cuando no hay nadie en espera. Convierte turnos vacíos en facturación."
    ),
    (
        "pedido_listo_retiro_v1",
        "UTILITY (Utilidad)",
        "Ópticas, Talleres Mecánicos, Imprentas, Reparaciones Técnicas",
        "«Hola {{1}}, te avisamos desde {{2}} que tu {{3}} ya se encuentra listo para retirar en {{4}}. Horario de atención: {{5}}.»",
        "• [👍 VOY HOY]\n• [💬 CONSULTAR]",
        "Acelera el cobro de saldos pendientes y despeja depósitos de productos listos."
    ),
    (
        "aviso_vencimiento_cuota_v1",
        "UTILITY (Utilidad)",
        "Gimnasios, Crossfit, Danza, Natatorios, Clubes, Academias",
        "«Hola {{1}}, te recordamos de {{2}} que tu cuota de {{3}} vence el {{4}}. Podés abonarla por recepción o mediante el siguiente botón seguro:»",
        "• [💳 PAGAR CUOTA]\n(Botón URL dinámico a Mercado Pago)",
        "Cobranza cordial automática del 1 al 10. Erradica la morosidad sin discusiones."
    ),
    (
        "campana_promo_v1",
        "MARKETING",
        "Reactivación de Ex-Socios en Gyms, Promociones de Temporada, B2B",
        "«Hola {{1}}! En {{2}} tenemos una novedad especial para vos: {{3}}. ¿Te gustaría aprovecharla?»",
        "• [🔥 QUIERO LA PROMO]\n• [❌ NO POR AHORA]",
        "Recupera clientes que no van hace 60+ días. Notifica al dueño en cuanto hay interés."
    )
]

for r in rows_sheet2:
    ws2.append(list(r))
    curr_row = ws2.max_row
    ws2.row_dimensions[curr_row].height = 45
    for c_idx, cell in enumerate(ws2[curr_row], 1):
        cell.font = normal_font
        cell.border = thin_border
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if c_idx == 1:
            cell.font = bold_font

auto_fit_columns(ws2)

# ==============================================================================
# HOJA 3: FORMATO EXCEL PARA EL ONBOARDING
# ==============================================================================
ws3 = wb.create_sheet(title="Checklist Onboarding y Datos")
ws3.views.sheetView[0].showGridLines = True

ws3.append(["ESTRUCTURA DE DATOS A PEDIR AL COMERCIO (PUESTA EN MARCHA ÁGIL)"])
ws3.merge_cells("A1:E1")
ws3["A1"].font = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")
ws3.row_dimensions[1].height = 28

ws3.append(["Modelo exacto de planilla que le solicitamos al cliente para cargar su base inicial en Sofía."])
ws3.merge_cells("A2:E2")
ws3["A2"].font = italic_font
ws3.row_dimensions[2].height = 18
ws3.append([])

headers3 = ["Columna en el Excel", "¿Obligatorio?", "Ejemplo de Contenido", "¿Para qué lo usa Sofía?", "Impacto en el Negocio"]
ws3.append(headers3)
style_header(ws3[4])
ws3.row_dimensions[4].height = 25

rows_sheet3 = [
    (
        "Nombre Completo",
        "SÍ",
        "María Gómez / Lucas Martínez",
        "Personalización de mensajes y cordialidad en el trato.",
        "Genera confianza inmediata; el paciente se siente atendido de forma VIP."
    ),
    (
        "Teléfono Celular",
        "SÍ",
        "5493434556677 (con código de área)",
        "Identificador único del paciente/cliente en el sistema.",
        "Permite que Sofía lo reconozca automáticamente sin menús ni hashtags."
    ),
    (
        "Próximo Turno / Vencimiento",
        "RECOMENDADO",
        "Jueves 08/10 a las 16:30 hs / 10/10/2026",
        "Dispara los recordatorios de 48h y 24h desde el día 1.",
        "Salva turnos y cobra cuotas de la misma semana de contratación."
    ),
    (
        "Profesional / Especialidad / Plan",
        "RECOMENDADO",
        "Dra. Silvina Gómez (Ortodoncia) / Pase Libre",
        "Segmenta agendas independientes y cotizaciones.",
        "Evita mezclar turnos entre médicos y profesionales distintos."
    ),
    (
        "Estado o Notas",
        "OPCIONAL",
        "Activo / Debe certificado médico / Prefiere tarde",
        "Contexto inyectado en el prompt de la IA de Sofía.",
        "Sofía puede recordarle requisitos especiales ('Recordá traer los estudios')."
    )
]

for r in rows_sheet3:
    ws3.append(list(r))
    curr_row = ws3.max_row
    ws3.row_dimensions[curr_row].height = 35
    for c_idx, cell in enumerate(ws3[curr_row], 1):
        cell.font = normal_font
        cell.border = thin_border
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if c_idx == 1:
            cell.font = bold_font

auto_fit_columns(ws3)

# ==============================================================================
# HOJA 4: CLÍNICAS Y MULTI-SECRETARIAS
# ==============================================================================
ws4 = wb.create_sheet(title="Multi-Secretarias y Equipos")
ws4.views.sheetView[0].showGridLines = True

ws4.append(["OPERATORIA PARA CLÍNICAS CON MÚLTIPLES SECRETARIAS Y TURNOS ROTATIVOS"])
ws4.merge_cells("A1:D1")
ws4["A1"].font = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")
ws4.row_dimensions[1].height = 28

ws4.append(["Cómo coordina Sofía la atención sin fricciones cuando hay varios puestos de recepción o turnos mañana/tarde."])
ws4.merge_cells("A2:D2")
ws4["A2"].font = italic_font
ws4.row_dimensions[2].height = 18
ws4.append([])

headers4 = ["Módulo / Función", "¿Cómo se Implementa?", "Ventaja Operativa para la Clínica", "Ejemplo de Uso Diario"]
ws4.append(headers4)
style_header(ws4[4])
ws4.row_dimensions[4].height = 25

rows_sheet4 = [
    (
        "Grupo de WhatsApp Interno ('Recepción Alvear')",
        "Sofía se suma como miembro del grupo de WhatsApp donde están las secretarias y la supervisora.",
        "Todas ven las cancelaciones y los avisos de las 15:00 hs en un solo lugar sin desfasajes de turno.",
        "A las 15:00 hs Sofía manda el resumen. La secretaria de la tarde escribe 'mantener a Carlos' en el grupo."
    ),
    (
        "Múltiples Líneas Autorizadas (Comandos de Voz)",
        "Se registran los celulares de cada secretaria (ej. Laura y Micaela) en la base de datos de la clínica.",
        "Cualquiera de las secretarias puede consultar o dar órdenes por audio desde su propio teléfono.",
        "Laura manda audio: 'Sofi, anotá sobreturno a Pedro mañana a las 11hs con Dr. Alvear'. Sofía lo agenda."
    ),
    (
        "Agendas Independientes por Doctor",
        "Cada profesional tiene su propia agenda, duración de consulta y lista de espera separada.",
        "Si se libera un turno de Odontología, no molesta a los pacientes de Ortodoncia ni kinesiología.",
        "El Rellena-Huecos y el Adelanta-Turnos solo contactan a pacientes de ESE doctor específico."
    ),
    (
        "Consultas de Estado de Agenda en Vivo",
        "Las secretarias pueden consultar los baches vacíos de mañana en cualquier momento.",
        "Ahorra tiempo en mostrador: no tienen que abrir planillas pesadas si están atendiendo la ventanilla.",
        "Secretaria escribe: '¿Quedó algún hueco mañana?'. Sofía devuelve la lista formateada en 2 segundos."
    )
]

for r in rows_sheet4:
    ws4.append(list(r))
    curr_row = ws4.max_row
    ws4.row_dimensions[curr_row].height = 40
    for c_idx, cell in enumerate(ws4[curr_row], 1):
        cell.font = normal_font
        cell.border = thin_border
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if c_idx == 1:
            cell.font = bold_font

auto_fit_columns(ws4)

# ==============================================================================
# HOJA 5: CRONOLOGÍA ANTI-HUECOS
# ==============================================================================
ws5 = wb.create_sheet(title="Cronología Anti-Huecos")
ws5.views.sheetView[0].showGridLines = True

ws5.append(["CRONOLOGÍA COMPLETA DEL SISTEMA ANTI-AUSENTISMO Y ADELANTA-TURNOS"])
ws5.merge_cells("A1:D1")
ws5["A1"].font = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")
ws5.row_dimensions[1].height = 28

ws5.append(["Línea de tiempo paso a paso desde 48 horas antes hasta la ocupación total de los consultorios."])
ws5.merge_cells("A2:D2")
ws5["A2"].font = italic_font
ws5.row_dimensions[2].height = 18
ws5.append([])

headers5 = ["Momento / Horario", "Acción Automática de Sofía", "¿Qué ve el Paciente o la Secretaria?", "Lógica y Objetivo"]
ws5.append(headers5)
style_header(ws5[4])
ws5.row_dimensions[4].height = 25

rows_sheet5 = [
    (
        "48 Horas Antes (Cálido y Natural)",
        "Sofía envía recordatorio cordial en texto plano sin botones invasivos.",
        "Paciente recibe: 'Hola Juan, te recordamos que pasado mañana tenés consulta...'.",
        "Aviso temprano sin presión. Si el paciente saluda o consulta algo, la IA responde con calidez."
    ),
    (
        "24 Horas Antes — 09:00 AM (Decisivo)",
        "Sofía envía mensaje con botones interactivos [SÍ, CONFIRMO] y [NO, REPROGRAMAR]. Plazo: 18:00 hs.",
        "Paciente recibe el mensaje oficial con botones y sabe que tiene casi todo el día para responder.",
        "Exige definición antes del final de la tarde para dar tiempo a cubrir cualquier cancelación."
    ),
    (
        "15:00 hs — Escudo de la Secretaria",
        "Sofía envía reporte confidencial a secretaría con los turnos de mañana aún sin confirmar (3 hs de margen).",
        "Secretaria recibe la lista. Puede llamar a pacientes conocidos o escribir: 'mantener a [Nombre]'.",
        "Le da el control absoluto al ser humano antes de cualquier corte automático."
    ),
    (
        "18:00 hs — Corte Automático",
        "Si no hubo confirmación ni intervención de secretaría, Sofía cancela con aviso de no concurrencia.",
        "El paciente sin confirmar recibe: 'Tu turno fue liberado. Por favor no concurras sin reprogramar'.",
        "Evita que el médico esté esperando en vano y libera el hueco para cubrirlo."
    ),
    (
        "18:01 hs — Rellena-Huecos (Lista de Espera)",
        "Sofía ofrece el turno liberado al Candidato #1 en lista de espera de a 1 persona por vez con 2 horas de margen.",
        "Candidato #1 recibe botones [TOMAR TURNO] o [NO, PASO]. Si toma, se agenda de inmediato.",
        "2 horas de margen es tiempo suficiente para responder. Si no responde o rechaza, pasa al Candidato #2."
    ),
    (
        "18:02 hs — Adelanta-Turnos (Fast-Track)",
        "Si la lista de espera está vacía, Sofía busca pacientes con turnos futuros (2 a 14 días) y les ofrece adelantar a mañana.",
        "Paciente del próximo martes recibe: 'Se liberó un turno para mañana, ¿te gustaría adelantarlo?'.",
        "Convierte un hueco muerto de mañana en un turno facturado, y libera el turno lejano para nuevos pacientes."
    )
]

for r in rows_sheet5:
    ws5.append(list(r))
    curr_row = ws5.max_row
    ws5.row_dimensions[curr_row].height = 42
    for c_idx, cell in enumerate(ws5[curr_row], 1):
        cell.font = normal_font
        cell.border = thin_border
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if c_idx == 1:
            cell.font = bold_font

auto_fit_columns(ws5)

output_path = "docs/Machete_Operativo_Comercial_Sofia.xlsx"
wb.save(output_path)
print(f"✅ Excel generado con éxito en: {output_path}")
