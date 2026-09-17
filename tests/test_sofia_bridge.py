import os
import tempfile
import pytest
import openpyxl
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.database import Base
from app.models.prospect import BridgeCommand, MerchantProduct
from app.services import excel_bridge
from tools.sofia_bridge import ExcelOperator
from main import app

@pytest.fixture
def db_session():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(bind=engine)
    Session = sessionmaker(bind=engine)
    db = Session()
    yield db
    db.close()

def test_excel_bridge_command_lifecycle(db_session):
    phone = "5493434991122"
    # 1. Enqueue command
    cmd = excel_bridge.enqueue_bridge_command(
        db=db_session,
        merchant_phone=phone,
        action="update_product",
        payload={"search": "cal loma negra", "updates": {"precio": 4800, "stock": 20}},
        sheet_name="Precios"
    )
    assert cmd.command_id is not None
    assert cmd.status == "pending"

    # 2. Retrieve pending commands
    pending = excel_bridge.get_pending_commands(db_session, phone)
    assert len(pending) == 1
    assert pending[0]["command_id"] == cmd.command_id
    assert pending[0]["payload"]["updates"]["precio"] == 4800

    # 3. Acknowledge command
    ack = excel_bridge.acknowledge_command(
        db=db_session,
        command_id=cmd.command_id,
        status="completed",
        result_message="Fila #2 actualizada con éxito."
    )
    assert ack is not None
    assert ack.status == "completed"
    assert "Fila #2" in ack.result_message

    # 4. Verify no more pending
    pending_after = excel_bridge.get_pending_commands(db_session, phone)
    assert len(pending_after) == 0

def test_pocket_price_lookup_and_formatting(db_session):
    phone = "5493434991122"
    # Seed a product
    p = MerchantProduct(
        merchant_phone=phone,
        supplier_name="Distribuidora Alem",
        name="Caño PVC 110mm Reforzado Tigre",
        cost_price=10000.0,
        price=14000.0,
        in_stock=True
    )
    db_session.add(p)
    db_session.commit()

    # Search by partial voice term
    res = excel_bridge.lookup_merchant_product(db_session, phone, "caño 110 tigre")
    assert res is not None
    assert res["name"] == "Caño PVC 110mm Reforzado Tigre"
    assert res["sale_price"] == 14000.0
    assert res["cost_price"] == 10000.0

    # Format text for WhatsApp / audio
    txt = excel_bridge.format_pocket_price_response(res)
    assert "Caño PVC 110mm Reforzado Tigre" in txt
    assert "$14.000" in txt
    assert "Distribuidora Alem" in txt
    assert "En stock ✅" in txt

def test_parse_pocket_price_query():
    # Various Argentine colloquial inquiries
    assert excel_bridge.parse_pocket_price_query("¿A cuánto tengo la cal Loma Negra?") is not None
    assert "cal loma negra" in excel_bridge.parse_pocket_price_query("a cuanto tengo la cal loma negra").lower()
    assert "caño 110" in excel_bridge.parse_pocket_price_query("¿Cuánto sale el caño 110?").lower()
    assert "cable 2.5" in excel_bridge.parse_pocket_price_query("precio de cable 2.5").lower()
    assert "lija al agua" in excel_bridge.parse_pocket_price_query("a cuánto cobro la lija al agua").lower()

    # Exclusions
    assert excel_bridge.parse_pocket_price_query("pasame la lista de precios") is None
    assert excel_bridge.parse_pocket_price_query("mandale el pedido a Carlos") is None

def test_excel_operator_local_execution():
    with tempfile.TemporaryDirectory() as tmpdir:
        excel_path = os.path.join(tmpdir, "ferreteria_test.xlsx")
        operator = ExcelOperator(excel_path)
        assert os.path.exists(excel_path)

        # 1. Update existing product in 'Precios' sheet
        res_update = operator.update_product(
            search_term="cal loma negra",
            updates={"precio": 5200, "stock": 12},
            sheet_name="Precios"
        )
        assert "actualizada" in res_update

        # Verify in file
        wb = openpyxl.load_workbook(excel_path)
        ws = wb["Precios"]
        # Find row for cal
        found = False
        for row in range(2, ws.max_row + 1):
            if "cal" in str(ws.cell(row=row, column=2).value).lower():
                assert ws.cell(row=row, column=4).value == 5200
                assert ws.cell(row=row, column=5).value == 12
                found = True
                break
        assert found is True

        # 2. Append sale to 'Ventas' sheet
        res_append = operator.append_row(
            values=["14/09", "19:40", "2 curvas 110 + 1 pegamento", 11600],
            sheet_name="Ventas"
        )
        assert "agregada" in res_append
        wb = openpyxl.load_workbook(excel_path)
        ws_v = wb["Ventas"]
        assert ws_v.max_row >= 2
        assert ws_v.cell(row=ws_v.max_row, column=3).value == "2 curvas 110 + 1 pegamento"

def test_bridge_api_endpoints():
    client = TestClient(app)

    # 1. Enqueue command via API
    resp = client.post("/api/v1/bridge/enqueue", json={
        "merchant_phone": "5491122334455",
        "action": "update_product",
        "payload": {"search": "cinta aisladora", "updates": {"precio": 1600}},
        "sheet_name": "Precios"
    })
    assert resp.status_code == 200
    cmd_id = resp.json()["command_id"]

    # 2. Get pending commands
    resp_get = client.get("/api/v1/bridge/commands", params={"merchant_phone": "5491122334455"})
    assert resp_get.status_code == 200
    cmds = resp_get.json()["commands"]
    assert any(c["command_id"] == cmd_id for c in cmds)

    # 3. Acknowledge command
    resp_ack = client.post("/api/v1/bridge/ack", json={
        "command_id": cmd_id,
        "status": "completed",
        "result_message": "Celda actualizada a $1.600"
    })
    assert resp_ack.status_code == 200
    assert resp_ack.json()["new_status"] == "completed"

    # 4. Test ZIP installer download
    resp_dl = client.get("/api/v1/bridge/download?merchant_phone=5491122334455")
    assert resp_dl.status_code == 200
    assert "application/zip" in resp_dl.headers.get("content-type", "")
    assert len(resp_dl.content) > 200

@pytest.mark.asyncio
async def test_boss_mode_bridge_and_pocket_price_integration(db_session):
    from app.services.boss_mode import process_boss_message

    phone = "5493434991122"
    # Seed product in DB
    p = MerchantProduct(
        merchant_phone=phone,
        supplier_name="Distribuidora San Martín",
        name="Lija al agua grano 100",
        cost_price=500.0,
        price=750.0,
        in_stock=True
    )
    db_session.add(p)
    db_session.commit()

    # 1. Pocket price voice query
    handled, reply, action = await process_boss_message(
        db=db_session,
        sender_phone=phone,
        text="Sofi, ¿a cuánto tengo la lija al agua?"
    )
    assert handled is True
    assert action == "pocket_price_query"
    assert "$750" in reply
    assert "Lija al agua" in reply

    # 2. Bridge Excel command
    handled_b, reply_b, action_b = await process_boss_message(
        db=db_session,
        sender_phone=phone,
        text="Sofi, cambiá en el excel el precio de la cal a 4800"
    )
    assert handled_b is True
    assert action_b == "bridge_excel_command"
    assert "SOFÍA BRIDGE EXCEL" in reply_b
    assert "actualiza en vivo" in reply_b

def test_bridge_ping_heartbeat_and_byok(db_session):
    client = TestClient(app)
    # 1. Ping with matching version and Gemini BYOK key
    resp = client.post("/api/v1/bridge/ping", json={
        "merchant_phone": "5493434991122",
        "version": "1.1.0",
        "gemini_api_key": "AIzaSyFakeKeyTest123"
    })
    assert resp.status_code == 200
    data = resp.json()
    assert data["status"] == "online"
    assert data["server_version"] == "1.1.0"
    assert data["client_version"] == "1.1.0"
    assert data["update_available"] is False
    assert data["merchant_phone"] == "5493434991122"

    # 2. Ping with older version -> update_available True
    resp_old = client.post("/api/v1/bridge/ping", json={
        "merchant_phone": "5493434991122",
        "version": "1.0.0"
    })
    assert resp_old.status_code == 200
    assert resp_old.json()["update_available"] is True

def test_detect_column_role_abbreviations():
    from tools.sofia_bridge import detect_column_role

    # Cost abbreviations
    assert detect_column_role("P.C.") == "cost"
    assert detect_column_role("P.C") == "cost"
    assert detect_column_role("p/c") == "cost"
    assert detect_column_role("COST") == "cost"
    assert detect_column_role("P. COSTO") == "cost"
    assert detect_column_role("PRECIO DE COSTO") == "cost"
    assert detect_column_role("Costo Unitario") == "cost"

    # Price abbreviations
    assert detect_column_role("P.V.") == "price"
    assert detect_column_role("P.V") == "price"
    assert detect_column_role("p/v") == "price"
    assert detect_column_role("PVP") == "price"
    assert detect_column_role("P. VENTA") == "price"
    assert detect_column_role("PRECIO VENTA") == "price"
    assert detect_column_role("PRECIO AL PUBLICO") == "price"
    assert detect_column_role("LISTA 1") == "price"
    assert detect_column_role("$") == "price"

    # Name / Detail abbreviations
    assert detect_column_role("ART.") == "name"
    assert detect_column_role("DESC.") == "name"
    assert detect_column_role("Detalle") == "name"
    assert detect_column_role("Mercadería") == "name"

    # Stock abbreviations
    assert detect_column_role("STK") == "stock"
    assert detect_column_role("CANT.") == "stock"
    assert detect_column_role("DISP.") == "stock"
    assert detect_column_role("Existencias") == "stock"

    # Code abbreviations
    assert detect_column_role("COD.") == "code"
    assert detect_column_role("SKU") == "code"
    assert detect_column_role("EAN") == "code"

def test_excel_operator_chaotic_format():
    """Prueba que el conector soporte planillas con títulos arriba, filas vacías y siglas raras."""
    with tempfile.TemporaryDirectory() as tmpdir:
        excel_path = os.path.join(tmpdir, "lista_desordenada.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Lista General de Precios"

        # Fila 1: Título o logo del local
        ws.cell(row=1, column=1, value="FERRETERIA SAN MIGUEL S.R.L. — SUCURSAL 1")
        # Fila 2: Subtítulo
        ws.cell(row=2, column=1, value="Lista actualizada al 15/09/2026")
        # Fila 3: Vacía

        # Fila 4: Encabezados reales con abreviaturas callejeras
        headers = ["COD.", "ART.", "P.C.", "P.V.", "STK"]
        for c, h in enumerate(headers, start=1):
            ws.cell(row=4, column=c, value=h)

        # Filas 5+: Productos
        ws.cell(row=5, column=1, value="T-10")
        ws.cell(row=5, column=2, value="Tornillo Fix 4x40mm (Caja x 100)")
        ws.cell(row=5, column=3, value=1200)
        ws.cell(row=5, column=4, value=1800)
        ws.cell(row=5, column=5, value=20)

        ws.cell(row=6, column=1, value="A-01")
        ws.cell(row=6, column=2, value="Alambre Galvanizado 100m")
        ws.cell(row=6, column=3, value=8000)
        ws.cell(row=6, column=4, value=12000)
        ws.cell(row=6, column=5, value=5)

        wb.save(excel_path)

        # Inicializar operador en este archivo no convencional
        op = ExcelOperator(excel_path)

        # 1. Modificar precio y stock por voz
        res = op.update_product(
            search_term="tornillo fix 4x40",
            updates={"precio": 2200, "stock": 45}
        )
        assert "actualizada" in res
        assert "Fila #5" in res

        # Verificar que efectivamente se escribió en la Fila 5, en las columnas 4 y 5
        wb_check = openpyxl.load_workbook(excel_path)
        ws_check = wb_check["Lista General de Precios"]
        assert ws_check.cell(row=5, column=4).value == 2200
        assert ws_check.cell(row=5, column=5).value == 45

        # 2. Agregar venta (debe autogenerar la hoja 'Ventas' porque no existía)
        res_v = op.append_row(["16/09", "18:50", "2 cajas tornillos", 4400])
        assert "agregada con éxito" in res_v
        wb_check2 = openpyxl.load_workbook(excel_path)
        assert "Ventas" in wb_check2.sheetnames
        ws_v = wb_check2["Ventas"]
        assert ws_v.cell(row=2, column=3).value == "2 cajas tornillos"

