import pytest
import io
import os
import json
import openpyxl
from unittest.mock import patch, AsyncMock
from app.config.settings import settings
from app.models.prospect import Prospect
from tests.conftest import TestingSessionLocal
from app.services.boss_mode import (
    process_boss_message,
    load_supplier_drafts,
    clear_supplier_draft,
    get_supplier_draft,
    get_client_faq_text,
    get_client_manual_text,
    parse_supplier_deletion_intent,
    parse_client_deletion_intent,
    parse_supplier_phone_update_intent
)
from app.models.prospect import SupplierDraftOrder, MerchantProduct
from app.services.catalog import catalog_service, ProductItem

@pytest.fixture
def db():
    session = TestingSessionLocal()
    try:
        yield session
    finally:
        session.query(Prospect).delete()
        session.commit()
        session.close()

@pytest.mark.asyncio
async def test_supplier_registration_and_listing(db):
    with patch("app.services.whatsapp.send_whatsapp_template", new_callable=AsyncMock), \
         patch("app.services.whatsapp.send_whatsapp_message", new_callable=AsyncMock):
        # 1. Register supplier
        handled, reply, action = await process_boss_message(
            db,
            settings.WHATSAPP_ALERT_PHONE,
            "Sofi, agendá al proveedor Bulonera del Litoral al 3434536447 que nos vende tornillos"
        )
        assert handled is True
        assert action == "supplier_registered"
        assert "PROVEEDOR REGISTRADO" in reply
        assert "Bulonera del Litoral" in reply
        assert "5493434536447" in reply

        # 2. List suppliers
        handled_l, reply_l, action_l = await process_boss_message(
            db,
            settings.WHATSAPP_ALERT_PHONE,
            "Sofi, cuáles son mis proveedores?"
        )
        assert handled_l is True
        assert action_l == "suppliers_list"
        assert "Bulonera del Litoral" in reply_l
        assert "5493434536447" in reply_l


@pytest.mark.asyncio
async def test_supplier_baskets_and_inquiries(db):
    clear_supplier_draft("Bulonera del Litoral")
    clear_supplier_draft("Pinturas Litoral")

    # 1. Add items for Bulonera
    handled1, reply1, action1 = await process_boss_message(
        db,
        settings.WHATSAPP_ALERT_PHONE,
        "Sofi, anotá para la Bulonera del Litoral 5 cajas de tornillos T1 y 2 alicates"
    )
    assert handled1 is True
    assert action1 == "basket_item_added"
    assert "Anotado" in reply1
    assert "Bulonera del Litoral" in reply1

    # 2. Add items for Pinturas Litoral
    handled2, reply2, action2 = await process_boss_message(
        db,
        settings.WHATSAPP_ALERT_PHONE,
        "Sofi, para Pinturas Litoral anotame 3 latas de látex blanco"
    )
    assert handled2 is True
    assert action2 == "basket_item_added"
    assert "Anotado" in reply2
    assert "Pinturas Litoral" in reply2

    # 3. Check specific basket for Bulonera
    handled_b, reply_b, action_b = await process_boss_message(
        db,
        settings.WHATSAPP_ALERT_PHONE,
        "Sofi, qué tengo para pedirle a la Bulonera del Litoral?"
    )
    assert handled_b is True
    assert action_b == "single_basket_detail"
    assert "BULONERA DEL LITORAL" in reply_b
    assert "tornillos" in reply_b.lower()

    # 4. Check all open baskets
    handled_all, reply_all, action_all = await process_boss_message(
        db,
        settings.WHATSAPP_ALERT_PHONE,
        "Sofi, pedidos pendientes a proveedores"
    )
    assert handled_all is True
    assert action_all == "all_baskets_summary"
    assert "Bulonera del Litoral" in reply_all
    assert "Pinturas Litoral" in reply_all


@pytest.mark.asyncio
async def test_supplier_basket_dispatch_and_clearing(db):
    clear_supplier_draft("Bulonera del Litoral")
    clear_supplier_draft("La Bulonera del Litoral")

    with patch("app.services.whatsapp.send_whatsapp_template", new_callable=AsyncMock) as mock_tpl, \
         patch("app.services.whatsapp.send_whatsapp_message", new_callable=AsyncMock) as mock_msg, \
         patch("app.services.whatsapp.send_whatsapp_document", new_callable=AsyncMock) as mock_doc:
        
        # Setup Bulonera supplier in DB
        await process_boss_message(
            db,
            settings.WHATSAPP_ALERT_PHONE,
            "Sofi, agendá al proveedor Bulonera del Litoral al 3434536447"
        )
        # Add items to basket
        await process_boss_message(
            db,
            settings.WHATSAPP_ALERT_PHONE,
            "Sofi, anotá para la Bulonera del Litoral 10 cajas de tornillos y 4 pinzas"
        )

        # Dispatch order using the basket items
        handled, reply, action = await process_boss_message(
            db,
            settings.WHATSAPP_ALERT_PHONE,
            "Sofi, mandale el pedido a Bulonera del Litoral"
        )
        assert handled is True
        assert action == "kiosk_order_dispatched"
        assert "¡Pedido despachado con éxito!" in reply
        assert "pasados en limpio" in reply
        assert "5493434536447" in reply

        mock_msg.assert_called()
        # Find the order dispatch message
        order_msgs = [c[1] for c in mock_msg.call_args_list if c[1]["to_phone"] == "5493434536447" and ("pedido formal" in c[1]["text"].lower() or "remito" in c[1]["text"].lower())]
        assert len(order_msgs) >= 1
        mock_doc.assert_called_once()
        assert mock_doc.call_args[1]["to_phone"] == "5493434536447"

        # Verify basket is now empty
        basket_after = get_supplier_draft("Bulonera del Litoral")
        assert basket_after is None


def test_catalog_supplier_tagging_and_excel_export():
    catalog_service.set_rubro("ferreteria")
    
    # Create sample supplier Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lista Precios"
    ws.append(["Código", "Descripción", "Precio"])
    ws.append(["FER-001", "Martillo Galponero 20mm", 19500])
    ws.append(["BUL-999", "Tornillos Autoperforantes 100u", 8500])
    
    buf = io.BytesIO()
    wb.save(buf)
    content = buf.getvalue()

    result = catalog_service.update_from_supplier_excel(
        content,
        filename="lista_bulonera_marzo.xlsx",
        supplier_name="Bulonera del Litoral"
    )
    assert result["status"] == "success"
    assert "Bulonera del Litoral" in result["whatsapp_message"]

    # Verify supplier is tagged in catalog
    matched_prod = catalog_service.find_product_exact_or_best("Martillo Galponero 20mm")
    assert matched_prod is not None
    assert matched_prod.price == 19500
    assert matched_prod.supplier == "Bulonera del Litoral"

    # Verify export includes Proveedor column
    export_path = "assets/test_export_suppliers.xlsx"
    catalog_service.export_to_excel(export_path)
    
    wb_exp = openpyxl.load_workbook(export_path)
    ws_exp = wb_exp.active
    headers = [cell.value for cell in ws_exp[1]]
    assert "Proveedor" in headers
    
    if os.path.exists(export_path):
        os.remove(export_path)


def test_merchant_client_dispatch_to_distribuidora_alem_simulation(db):
    from fastapi.testclient import TestClient
    from main import app
    from app.database import get_db
    from app.services.boss_mode import LAST_ONBOARDED_CLIENT

    app.dependency_overrides[get_db] = lambda: db
    client = TestClient(app)

    # 1. Setup Distribuidora Alem with Javier's phone in DB
    alem = Prospect(
        name="Distribuidora Alem",
        contact_name="Javier (Distribuidora Alem)",
        phone="5493434536447",
        city="Paraná",
        campaign="supplier",
        status="supplier"
    )
    db.add(alem)

    # 2. Setup Ferretería Nogoyá as the active merchant client
    ferreteria = Prospect(
        name="Ferretería Nogoyá",
        contact_name="Ricardo",
        phone="5493435112233",
        city="Nogoyá",
        campaign="client_onboarding",
        status="in_conversation"
    )
    db.add(ferreteria)
    db.commit()

    LAST_ONBOARDED_CLIENT.clear()
    LAST_ONBOARDED_CLIENT.update({
        "phone": "5493435112233",
        "business_name": "Ferretería Nogoyá",
        "contact_name": "Ricardo",
        "business_type": "ferreteria",
        "city": "Nogoyá"
    })

    with patch("app.services.whatsapp.send_whatsapp_template", new_callable=AsyncMock) as mock_tpl, \
         patch("app.services.whatsapp.send_whatsapp_message", new_callable=AsyncMock) as mock_msg, \
         patch("app.services.whatsapp.send_whatsapp_document", new_callable=AsyncMock) as mock_doc:
        
        # Ricardo sends dispatch directive from his phone via webhook
        payload = {
            "phone": "5493435112233",
            "message": "Sofi, mandale el pedido a Distribuidora Alem con 4 martillos y 2 alicates",
            "complex_name": "Ferretería Nogoyá",
            "contact_name": "Ricardo",
            "city": "Nogoyá"
        }
        response = client.post("/webhook", json=payload)
        assert response.status_code == 200
        res = response.json()

        assert res["status"] == "success"
        assert res["action"] == "kiosk_order_dispatched"
        assert "Distribuidora Alem" in res["reply"]
        assert "5493434536447" in res["reply"]

        mock_msg.assert_called()
        dispatched_to_phones = [call[1]["to_phone"] for call in mock_msg.call_args_list]
        assert "5493434536447" in dispatched_to_phones
        assert "5493435112233" in dispatched_to_phones


def test_client_onboarding_5_options_and_quick_replies(db):
    from fastapi.testclient import TestClient
    from main import app
    client = TestClient(app)

    # 1. Setup client onboarding merchant
    merchant = Prospect(
        name="Ferretería El Triángulo",
        contact_name="Carlos",
        phone="5493436001122",
        city="Paraná",
        campaign="client_onboarding",
        status="in_conversation"
    )
    db.add(merchant)
    db.commit()

    with patch("app.services.whatsapp.send_whatsapp_message", new_callable=AsyncMock) as mock_msg, \
         patch("app.services.whatsapp.notify_owner_demo_requested", new_callable=AsyncMock) as mock_demo_notify:

        # Test A: Menu repetition with 5 options
        resp = client.post("/webhook", json={
            "phone": "5493436001122",
            "message": "Hola Sofi"
        })
        assert resp.status_code == 200
        data = resp.json()
        assert data.get("guided_menu") is True
        assert "catálogo de demostración" in data["reply"]
        assert "1️⃣" in data["reply"]
        assert "5️⃣" in data["reply"]
        assert "proveedores tengo registrados" in data["reply"]

        # Test B: Option 4 - How to upload / forward lists
        resp = client.post("/webhook", json={
            "phone": "5493436001122",
            "message": "4"
        })
        assert resp.status_code == 200
        data = resp.json()
        assert data.get("catalog_instructions") is True
        assert "PDF o Excel" in data["reply"]

        # Test C: Option 5 - Registered suppliers query
        resp = client.post("/webhook", json={
            "phone": "5493436001122",
            "message": "¿Qué proveedores tengo registrados?"
        })
        assert resp.status_code == 200
        data = resp.json()
        assert data.get("suppliers_summary") is True
        assert "PROVEEDORES Y LISTAS REGISTRADAS" in data["reply"]

        # Test D: Cold outreach quick reply - 'Ver demostración'
        lead = Prospect(
            name="Comercio Prospecto",
            contact_name="Martín",
            phone="5493436998877",
            city="Paraná",
            campaign="ai_agency",
            status="contacted"
        )
        db.add(lead)
        db.commit()

        resp = client.post("/webhook", json={
            "phone": "5493436998877",
            "message": "Ver demostración"
        })
        assert resp.status_code == 200
        data = resp.json()
        assert data.get("demo_requested") is True
        mock_demo_notify.assert_called_once()

        # Test E: Cold outreach quick reply - 'Ahora no, gracias'
        resp = client.post("/webhook", json={
            "phone": "5493436998877",
            "message": "Ahora no, gracias"
        })
        assert resp.status_code == 200
        data = resp.json()
        assert data.get("action") == "opt_out"


def test_merchant_client_register_supplier_with_auto_presentation(db):
    from fastapi.testclient import TestClient
    from main import app
    from app.database import get_db

    app.dependency_overrides[get_db] = lambda: db
    client = TestClient(app)

    # 1. Setup Kiosco Alameda as the merchant client
    kiosco = Prospect(
        name="Kiosco Alameda",
        contact_name="Marcelo",
        phone="5493435112233",
        city="Paraná",
        campaign="client_onboarding",
        status="in_conversation"
    )
    db.add(kiosco)
    db.commit()

    with patch("app.services.whatsapp.send_whatsapp_template", new_callable=AsyncMock) as mock_tpl, \
         patch("app.services.whatsapp.send_whatsapp_message", new_callable=AsyncMock) as mock_msg:
        mock_tpl.return_value = False

        # Marcelo instructs Sofia by text to register Carlos from Distribuidora El Progreso
        payload = {
            "phone": "5493435112233",
            "message": "Sofi, agendá al proveedor Distribuidora El Progreso al 3434536447"
        }
        resp = client.post("/webhook", json=payload)
        assert resp.status_code == 200
        data = resp.json()

        assert data.get("status") == "success"
        assert data.get("action") == "supplier_registered"
        assert "PROVEEDOR REGISTRADO Y CONTACTADO" in data["reply"]
        assert "Distribuidora El Progreso" in data["reply"]
        assert "5493434536447" in data["reply"]

        # Check DB
        sup_db = db.query(Prospect).filter(Prospect.phone == "5493434536447").first()
        assert sup_db is not None
        assert sup_db.name == "Distribuidora El Progreso"
        assert sup_db.business_type == "proveedor"
        assert sup_db.campaign == "supplier"

        # Check Meta Template dispatch to the supplier
        mock_tpl.assert_called_once()
        tpl_args = mock_tpl.call_args[1]
        assert tpl_args["to_phone"] == "5493434536447"
        assert tpl_args["template_name"] == "presentacion_proveedor_v1"

        # Check conversational message to supplier
        dispatched_messages = [call[1] for call in mock_msg.call_args_list if call[1]["to_phone"] == "5493434536447"]
        assert len(dispatched_messages) >= 1
        intro_text = dispatched_messages[0]["text"]
        assert "Marcelo de Kiosco Alameda" in intro_text
        assert "lista de precios" in intro_text
        assert "Agendá este contacto" in intro_text
        assert "Agendado" in intro_text


def test_meta_contacts_vcard_registration(db):
    from fastapi.testclient import TestClient
    from main import app
    from app.database import get_db

    app.dependency_overrides[get_db] = lambda: db
    client = TestClient(app)

    # Setup merchant
    ferreteria = Prospect(
        name="Ferretería Nogoyá",
        contact_name="Ricardo",
        phone="5493435112233",
        campaign="client_onboarding",
        status="in_conversation"
    )
    db.add(ferreteria)
    db.commit()

    with patch("app.services.whatsapp.send_whatsapp_template", new_callable=AsyncMock) as mock_tpl, \
         patch("app.services.whatsapp.send_whatsapp_message", new_callable=AsyncMock) as mock_msg:

        # Simulate Meta Cloud API incoming contact payload
        meta_payload = {
            "object": "whatsapp_business_account",
            "entry": [
                {
                    "id": "123456789",
                    "changes": [
                        {
                            "value": {
                                "messaging_product": "whatsapp",
                                "contacts": [{"profile": {"name": "Ricardo"}}],
                                "messages": [
                                    {
                                        "from": "5493435112233",
                                        "id": "wamid.contact.123",
                                        "timestamp": "1726000000",
                                        "type": "contacts",
                                        "contacts": [
                                            {
                                                "name": {
                                                    "first_name": "Carlos",
                                                    "formatted_name": "Carlos Molinos"
                                                },
                                                "phones": [
                                                    {
                                                        "phone": "+54 9 343 499-1122",
                                                        "type": "CELL"
                                                    }
                                                ]
                                            }
                                        ]
                                    }
                                ]
                            }
                        }
                    ]
                }
            ]
        }

        resp = client.post("/webhook", json=meta_payload)
        assert resp.status_code == 200
        data = resp.json()
        assert data.get("action") == "supplier_registered"

        # Check DB
        sup_db = db.query(Prospect).filter(Prospect.phone == "5493434991122").first()
        assert sup_db is not None
        assert sup_db.name in ["Carlos Molinos", "Molinos"]
        assert "Carlos" in (sup_db.contact_name or "")
        assert sup_db.business_type == "proveedor"


def test_client_manual_and_option_6_guidance(db):
    from fastapi.testclient import TestClient
    from main import app
    from app.database import get_db

    app.dependency_overrides[get_db] = lambda: db
    client = TestClient(app)

    kiosco = Prospect(
        name="Kiosco Alameda",
        contact_name="Marcelo",
        phone="5493435112233",
        campaign="client_onboarding",
        status="in_conversation"
    )
    db.add(kiosco)
    db.commit()

    with patch("app.services.whatsapp.send_whatsapp_message", new_callable=AsyncMock) as mock_msg:
        # 1. Ask for manual
        resp = client.post("/webhook", json={"phone": "5493435112233", "message": "manual"})
        assert resp.status_code == 200
        data = resp.json()
        assert data.get("client_manual_sent") is True
        assert "Agendar proveedores nuevos en 1 toque" in data["reply"]
        assert "Sofía - Compras" in data["reply"]

        # 2. Ask for option 6
        resp = client.post("/webhook", json={"phone": "5493435112233", "message": "6"})
        assert resp.status_code == 200
        data = resp.json()
        assert data.get("supplier_reg_guide") is True
        assert "CÓMO AGENDAR PROVEEDORES NUEVOS" in data["reply"]
        assert "Dictámelo por audio o texto" in data["reply"]
        assert "compartime su contacto de WhatsApp" in data["reply"]

        # 3. Menu greeting includes Option 6
        resp = client.post("/webhook", json={"phone": "5493435112233", "message": "hola"})
        assert resp.status_code == 200
        data = resp.json()
        assert data.get("guided_menu") is True
        assert "6️⃣" in data["reply"]
        assert "agendá al proveedor" in data["reply"].lower() or "agendá a" in data["reply"].lower()


@pytest.mark.asyncio
async def test_parse_and_apply_supplier_price_updates():
    from app.services.catalog import parse_supplier_price_update_text, catalog_service
    catalog_service.set_rubro("distribuidora")

    # Find initial price of Azúcar
    azucar = catalog_service.find_product_exact_or_best("Azúcar Ledesma")
    assert azucar is not None
    initial_p = azucar.price

    # 1. Test parsing of percentage and fixed price in single text
    msg = "Hola gente, aumentó el azúcar Ledesma un 5% y el aceite Cañuelas pasa a $2400"
    data = await parse_supplier_price_update_text(msg)
    assert data["is_price_update"] is True
    assert len(data["updates"]) == 2

    # 2. Process updates in catalog
    modified = catalog_service.process_supplier_price_updates(data["updates"])
    assert len(modified) >= 1

    # Verify Azúcar price was updated by 5%
    azucar_after = catalog_service.find_product_exact_or_best("Azúcar Ledesma")
    expected_new = round(initial_p * 1.05, 2)
    assert azucar_after.price == expected_new


def test_webhook_incoming_supplier_price_update_and_merchant_alert(db):
    from fastapi.testclient import TestClient
    from main import app
    from app.database import get_db

    app.dependency_overrides[get_db] = lambda: db
    client = TestClient(app)

    # 1. Setup store merchant
    kiosco = Prospect(
        name="Kiosco Alameda",
        contact_name="Marcelo",
        phone="5493435112233",
        campaign="client_onboarding",
        status="in_conversation"
    )
    db.add(kiosco)
    db.commit()

    # 2. Setup supplier linked to merchant
    sup_meta = {"merchant_phone": "5493435112233", "merchant_biz": "Kiosco Alameda"}
    distribuidora = Prospect(
        name="Distribuidora El Progreso",
        contact_name="Carlos",
        phone="5493434556677",
        business_type="proveedor",
        campaign="supplier",
        notes=json.dumps(sup_meta)
    )
    db.add(distribuidora)
    db.commit()

    with patch("app.services.whatsapp.send_whatsapp_message", new_callable=AsyncMock) as mock_msg:
        # Supplier sends price increase via WhatsApp
        payload = {
            "phone": "5493434556677",
            "message": "Hola Sofía, a partir del lunes el azúcar Ledesma sube un 5%"
        }
        resp = client.post("/webhook", json=payload)
        assert resp.status_code == 200
        data = resp.json()
        assert data.get("status") == "success"
        assert data.get("action") == "supplier_price_updated"
        assert "Ya registré los aumentos" in data["reply"]
        assert "Azúcar Ledesma" in data["reply"]

        # Check that merchant received alert
        merchant_alerts = [c[1] for c in mock_msg.call_args_list if c[1]["to_phone"] == "5493435112233"]
        assert len(merchant_alerts) >= 1
        assert "AVISO DE AUMENTO DE TU PROVEEDOR" in merchant_alerts[0]["text"]
        assert "Distribuidora El Progreso" in merchant_alerts[0]["text"]
        assert "Azúcar Ledesma" in merchant_alerts[0]["text"]


def test_webhook_incoming_supplier_acknowledgment(db):
    from fastapi.testclient import TestClient
    from main import app
    from app.database import get_db

    app.dependency_overrides[get_db] = lambda: db
    client = TestClient(app)

    # Setup merchant and supplier
    kiosco = Prospect(
        name="Kiosco Alameda",
        contact_name="Marcelo",
        phone="5493435112233",
        campaign="client_onboarding"
    )
    db.add(kiosco)
    db.commit()

    distribuidora = Prospect(
        name="Distribuidora El Progreso",
        contact_name="Carlos",
        phone="5493434556677",
        business_type="proveedor",
        campaign="supplier",
        notes=json.dumps({"merchant_phone": "5493435112233"})
    )
    db.add(distribuidora)
    db.commit()

    with patch("app.services.whatsapp.send_whatsapp_message", new_callable=AsyncMock) as mock_msg:
        # Supplier confirms with "Agendado"
        payload = {
            "phone": "5493434556677",
            "message": "Agendado"
        }
        resp = client.post("/webhook", json=payload)
        assert resp.status_code == 200
        data = resp.json()
        assert data.get("action") == "supplier_acknowledged"
        assert "Muchas gracias" in data["reply"]

        # Merchant gets notification that supplier confirmed
        merchant_alerts = [c[1] for c in mock_msg.call_args_list if c[1]["to_phone"] == "5493435112233"]
        assert len(merchant_alerts) >= 1
        assert "PROVEEDOR CONFIRMÓ RECEPCIÓN" in merchant_alerts[0]["text"]


@pytest.mark.asyncio
async def test_merchant_direct_price_update_directive(db):
    from app.services.catalog import catalog_service
    catalog_service.set_rubro("distribuidora")

    # Store merchant says: "Sofi, aumentó el azúcar Ledesma un 5%"
    handled, reply, action = await process_boss_message(
        db,
        settings.WHATSAPP_ALERT_PHONE,
        "Sofi, aumentó el azúcar Ledesma un 5%"
    )
    assert handled is True
    assert action == "merchant_price_update_applied"
    assert "CATÁLOGO ACTUALIZADO" in reply
    assert "Azúcar Ledesma" in reply


@pytest.mark.asyncio
async def test_free_dictation_and_7day_rule(db):
    clear_supplier_draft("Distribuidora Alem")
    clear_supplier_draft("Distribuidora Nogoyá")

    # Dictate items without specifying supplier
    handled, reply, action = await process_boss_message(
        db,
        settings.WHATSAPP_ALERT_PHONE,
        "Sofi, anotame 5 paquetes de harina 000 y 3 botellas de aceite"
    )
    assert handled is True
    assert action == "basket_item_added"
    assert "Anotado" in reply
    # Must specify street action prompts
    assert "Mostrame lo que le tengo anotado" in reply or "Mandale el pedido" in reply


@pytest.mark.asyncio
async def test_batch_executive_summary_and_street_language(db):
    clear_supplier_draft("Distribuidora Alem")
    clear_supplier_draft("Mayorista Central")

    # Merchant dictates 6 items (batch > 3)
    handled, reply, action = await process_boss_message(
        db,
        settings.WHATSAPP_ALERT_PHONE,
        "Sofi, anotá 10 cajas de alfajores, 5 de coca, 4 de fideos, 2 de aceite, 6 de pure de tomate y 3 de mayonesa"
    )
    assert handled is True
    assert action == "basket_item_added"
    # Should use Executive Summary grouping instead of endless chat
    assert "artículos repartidos por mejor precio" in reply or "artículos" in reply
    # Should check supplier grouping with freshness badge
    assert ("🟢" in reply or "⚠️" in reply)
    # Street prompts
    assert "Mostrame lo que le tengo anotado" in reply or "Mandale el pedido" in reply

    # Merchant inquires using Argentine street slang
    handled_q, reply_q, action_q = await process_boss_message(
        db,
        settings.WHATSAPP_ALERT_PHONE,
        "Sofi, qué le tengo anotado a Alem?"
    )
    assert handled_q is True
    assert action_q == "single_basket_detail"
    assert "LO QUE TENÉS ANOTADO PARA" in reply_q
    assert "canasta" not in reply_q.lower()


@pytest.mark.asyncio
async def test_client_faq_and_security_guide(db):
    # 1. Check content of get_client_faq_text
    faq = get_client_faq_text()
    assert len(faq) < 4000, f"FAQ text too long ({len(faq)} chars) for Meta 4096 character limit"
    assert "GUÍA DE SEGURIDAD COMERCIAL Y PREGUNTAS FRECUENTES" in faq
    assert "BLOQUE 1: PRECIOS, INFLACIÓN" in faq
    assert "BLOQUE 2: PROVEEDORES, VIAJANTES" in faq
    assert "BLOQUE 3: PRIVACIDAD, AUDIOS" in faq
    assert "Regla de los 7 días" in faq
    assert "viajante que viene a visitarme en persona" in faq
    assert "eliminar o dar de baja a un proveedor" in faq
    assert "confidencialidad es 100% estricta" in faq
    assert "circuito cerrado y profesional" in faq
    assert "únicamente se comunica con vos y con los distribuidores" in faq
    assert "Resumen Ejecutivo" in faq
    assert "JAMÁS!" in faq

    # 2. Check get_client_manual_text references 'dudas'
    manual = get_client_manual_text()
    assert "«dudas»" in manual

    # 3. Check view triggers for FAQ (clean, without artificial headers)
    triggers = ["dudas", "faq", "preguntas frecuentes", "que pasa si"]
    for t in triggers:
        handled, reply, action = await process_boss_message(
            db,
            settings.WHATSAPP_ALERT_PHONE,
            f"Sofi, {t}"
        )
        assert handled is True
        assert action == "boss_faq_view"
        assert "GUÍA DE SEGURIDAD COMERCIAL Y PREGUNTAS FRECUENTES" in reply
        assert "Listo para reenviar" not in reply
        assert "enviar dudas al" not in reply

    # 4. Check view trigger for manual (clean, without artificial headers)
    handled_m, reply_m, action_m = await process_boss_message(
        db,
        settings.WHATSAPP_ALERT_PHONE,
        "Sofi, manual"
    )
    assert handled_m is True
    assert action_m == "boss_manual_view"
    assert "TU CENTRAL DE COMPRAS EN WHATSAPP" in reply_m
    assert "Listo para reenviar" not in reply_m
    assert "enviar manual al" not in reply_m


@pytest.mark.asyncio
async def test_supplier_deletion_intent_and_execution(db):
    from app.services.boss_mode import add_items_to_supplier_draft

    # 1. Test intent parsing
    del1 = parse_supplier_deletion_intent("Sofi, eliminar proveedor Distribuidora San José")
    assert del1["is_supplier_deletion"] is True
    assert del1["supplier_name"] == "Distribuidora San José"

    del2 = parse_supplier_deletion_intent("borrar al proveedor Alem")
    assert del2["is_supplier_deletion"] is True
    assert del2["supplier_name"] == "Alem"

    del3 = parse_supplier_deletion_intent("dar de baja distribuidora Litoral")
    assert del3["is_supplier_deletion"] is True
    assert del3["supplier_name"] == "Litoral"

    del4 = parse_supplier_deletion_intent("Sofi, eliminar proveedor")
    assert del4["is_supplier_deletion"] is True
    assert del4["supplier_name"] is None

    del5 = parse_supplier_deletion_intent("hola que tal")
    assert del5["is_supplier_deletion"] is False

    # 2. Register supplier in DB and open draft
    sup_test = Prospect(
        name="Distribuidora San José",
        contact_name="José",
        phone="5493434998877",
        business_type="proveedor",
        campaign="supplier"
    )
    db.add(sup_test)
    db.commit()

    add_items_to_supplier_draft("Distribuidora San José", [{"product_name": "Tornillos", "quantity": 10, "unit_price": 500}])
    draft_before = get_supplier_draft("Distribuidora San José")
    assert draft_before is not None
    assert len(draft_before.get("items", [])) == 1

    # 3. Delete registered supplier
    handled, reply, action = await process_boss_message(
        db,
        settings.WHATSAPP_ALERT_PHONE,
        "Sofi, eliminar proveedor Distribuidora San José"
    )
    assert handled is True
    assert action == "supplier_deleted"
    assert "PROVEEDOR ELIMINADO CON ÉXITO" in reply
    assert "Distribuidora San José" in reply

    # Verify supplier is gone from DB
    deleted_in_db = db.query(Prospect).filter(Prospect.name == "Distribuidora San José").first()
    assert deleted_in_db is None

    # Verify draft is cleared
    draft_after = get_supplier_draft("Distribuidora San José")
    assert draft_after is None

    # 4. Deleting again gives supplier_not_found
    handled2, reply2, action2 = await process_boss_message(
        db,
        settings.WHATSAPP_ALERT_PHONE,
        "Sofi, eliminar proveedor Distribuidora San José"
    )
    assert handled2 is True
    assert action2 == "supplier_not_found"
    assert "No encontré ningún proveedor" in reply2

    # 5. Calling without supplier name gives supplier_delete_needs_name
    handled3, reply3, action3 = await process_boss_message(
        db,
        settings.WHATSAPP_ALERT_PHONE,
        "eliminar proveedor"
    )
    assert handled3 is True
    assert action3 == "supplier_delete_needs_name"
    assert "¿Qué proveedor te gustaría dar de baja?" in reply3


@pytest.mark.asyncio
async def test_webhook_merchant_client_faq_and_deletion(db):
    from fastapi.testclient import TestClient
    from main import app
    from app.database import get_db

    app.dependency_overrides[get_db] = lambda: db
    client = TestClient(app)

    client_lead = Prospect(
        name="Ferretería El Tornillo",
        contact_name="Marcos",
        phone="5493435112233",
        business_type="ferreteria",
        campaign="client_onboarding",
        status="in_conversation"
    )
    db.add(client_lead)
    db.commit()

    with patch("app.services.whatsapp.send_whatsapp_message", new_callable=AsyncMock):
        # 1. Merchant asks for 'dudas'
        res_faq = client.post("/webhook", json={"phone": "5493435112233", "message": "dudas"})
        assert res_faq.status_code == 200
        res_data = res_faq.json()
        assert res_data.get("client_faq_sent") is True
        assert "GUÍA DE SEGURIDAD COMERCIAL" in res_data.get("reply", "")

        # 2. Merchant sees menu with option 7
        res_menu = client.post("/webhook", json={"phone": "5493435112233", "message": "menu"})
        assert res_menu.status_code == 200
        assert "7️⃣ _Escribí «manual» para ver cómo usarme o «dudas»" in res_menu.json().get("reply", "")


@pytest.mark.asyncio
async def test_whatsapp_message_chunking_for_meta_limits():
    from app.services.whatsapp import split_whatsapp_message, send_whatsapp_message
    from unittest.mock import AsyncMock, patch

    # 1. Message under 3800 chars remains 1 chunk
    short_text = "Hola, este es un mensaje corto de prueba."
    chunks_short = split_whatsapp_message(short_text, max_chars=3800)
    assert len(chunks_short) == 1
    assert chunks_short[0] == short_text

    # 2. Huge message over 4000 chars is split cleanly
    para1 = "Párrafo 1 con información relevante. " * 50  # ~1900 chars
    para2 = "Párrafo 2 con más datos comerciales. " * 50   # ~1900 chars
    para3 = "Párrafo 3 con las conclusiones finales. " * 30  # ~1200 chars
    huge_text = f"{para1}\n\n{para2}\n\n{para3}"  # ~5000 chars

    chunks_huge = split_whatsapp_message(huge_text, max_chars=3800)
    assert len(chunks_huge) >= 2
    for c in chunks_huge:
        assert len(c) <= 3800

    # 3. send_whatsapp_message dispatches all chunks
    with patch("app.services.whatsapp._send_single_whatsapp_message", new_callable=AsyncMock) as mock_single:
        mock_single.return_value = True
        success = await send_whatsapp_message("5493434536447", huge_text)
        assert success is True
        assert mock_single.call_count == len(chunks_huge)


@pytest.mark.asyncio
async def test_supplier_price_update_broadcasts_to_all_linked_merchants(db):
    from fastapi.testclient import TestClient
    from main import app
    from app.database import get_db

    app.dependency_overrides[get_db] = lambda: db
    client = TestClient(app)

    m1_phone = "5493435111111"
    m2_phone = "5493435222222"
    supplier_phone = "5493435333333"

    # 1. Setup Merchant 1 and Merchant 2
    merchant_1 = Prospect(
        name="Ferretería Marcos",
        contact_name="Marcos",
        phone=m1_phone,
        business_type="ferreteria",
        campaign="client_onboarding",
        status="in_conversation"
    )
    merchant_2 = Prospect(
        name="Almacén Laura",
        contact_name="Laura",
        phone=m2_phone,
        business_type="almacen",
        campaign="client_onboarding",
        status="in_conversation"
    )
    # 2. Both merchants register Carlos (same supplier phone)
    sup_m1 = Prospect(
        merchant_phone=m1_phone,
        name="Distribuidora El Progreso",
        contact_name="Carlos",
        phone=supplier_phone,
        business_type="proveedor",
        campaign="supplier",
        notes=json.dumps({"merchant_phone": m1_phone})
    )
    sup_m2 = Prospect(
        merchant_phone=m2_phone,
        name="Distribuidora El Progreso",
        contact_name="Carlos",
        phone=supplier_phone,
        business_type="proveedor",
        campaign="supplier",
        notes=json.dumps({"merchant_phone": m2_phone})
    )
    db.add_all([merchant_1, merchant_2, sup_m1, sup_m2])
    db.commit()

    sent_messages = []

    async def mock_send(to_phone, text):
        sent_messages.append({"to_phone": to_phone, "text": text})
        return True

    with patch("app.services.whatsapp.send_whatsapp_message", side_effect=mock_send):
        # Carlos sends a price increase message
        res = client.post("/webhook", json={
            "phone": supplier_phone,
            "message": "Hola Sofía, a partir de mañana el aceite Cañuelas sube un 8% y la harina Pureza sube 5%"
        })
        assert res.status_code == 200
        res_data = res.json()
        assert res_data.get("action") == "supplier_price_updated"

        # Give background tasks a brief moment to run
        import asyncio
        await asyncio.sleep(0.05)

        # Assert Carlos received confirmation
        carlos_msgs = [m for m in sent_messages if m["to_phone"] == supplier_phone]
        assert len(carlos_msgs) >= 1
        assert "aumentos informados" in carlos_msgs[0]["text"].lower() or "registré" in carlos_msgs[0]["text"].lower()

        # Assert Merchant 1 received price alert
        m1_msgs = [m for m in sent_messages if m["to_phone"] == m1_phone]
        assert len(m1_msgs) >= 1
        assert "AVISO DE AUMENTO DE TU PROVEEDOR" in m1_msgs[0]["text"]
        assert "Distribuidora El Progreso" in m1_msgs[0]["text"]

        # Assert Merchant 2 ALSO received price alert (multi-merchant broadcast!)
        m2_msgs = [m for m in sent_messages if m["to_phone"] == m2_phone]
        assert len(m2_msgs) >= 1
        assert "AVISO DE AUMENTO DE TU PROVEEDOR" in m2_msgs[0]["text"]
        assert "Distribuidora El Progreso" in m2_msgs[0]["text"]


@pytest.mark.asyncio
async def test_supplier_ack_broadcasts_to_all_linked_merchants(db):
    from fastapi.testclient import TestClient
    from main import app
    from app.database import get_db

    app.dependency_overrides[get_db] = lambda: db
    client = TestClient(app)

    m1_phone = "5493435111111"
    m2_phone = "5493435222222"
    supplier_phone = "5493435333333"

    sup_m1 = Prospect(
        merchant_phone=m1_phone,
        name="Distribuidora El Progreso",
        contact_name="Carlos",
        phone=supplier_phone,
        business_type="proveedor",
        campaign="supplier"
    )
    sup_m2 = Prospect(
        merchant_phone=m2_phone,
        name="Distribuidora El Progreso",
        contact_name="Carlos",
        phone=supplier_phone,
        business_type="proveedor",
        campaign="supplier"
    )
    db.add_all([sup_m1, sup_m2])
    db.commit()

    sent_messages = []

    async def mock_send(to_phone, text):
        sent_messages.append({"to_phone": to_phone, "text": text})
        return True

    with patch("app.services.whatsapp.send_whatsapp_message", side_effect=mock_send):
        res = client.post("/webhook", json={
            "phone": supplier_phone,
            "message": "Agendado gracias"
        })
        assert res.status_code == 200
        assert res.json().get("action") == "supplier_acknowledged"

        import asyncio
        await asyncio.sleep(0.05)

        # Both merchants received acknowledgment alert
        m1_msgs = [m for m in sent_messages if m["to_phone"] == m1_phone]
        m2_msgs = [m for m in sent_messages if m["to_phone"] == m2_phone]
        assert len(m1_msgs) >= 1
        assert "PROVEEDOR CONFIRMÓ RECEPCIÓN" in m1_msgs[0]["text"]
        assert len(m2_msgs) >= 1
        assert "PROVEEDOR CONFIRMÓ RECEPCIÓN" in m2_msgs[0]["text"]


@pytest.mark.asyncio
async def test_supplier_excel_upload_broadcasts_to_all_linked_merchants(db):
    from fastapi.testclient import TestClient
    from main import app
    from app.database import get_db
    from app.models.prospect import MerchantProduct
    import base64

    app.dependency_overrides[get_db] = lambda: db
    client = TestClient(app)

    m1_phone = "5493435111111"
    m2_phone = "5493435222222"
    supplier_phone = "5493435333333"

    sup_m1 = Prospect(
        merchant_phone=m1_phone,
        name="Distribuidora El Progreso",
        contact_name="Carlos",
        phone=supplier_phone,
        business_type="proveedor",
        campaign="supplier"
    )
    sup_m2 = Prospect(
        merchant_phone=m2_phone,
        name="Distribuidora El Progreso",
        contact_name="Carlos",
        phone=supplier_phone,
        business_type="proveedor",
        campaign="supplier"
    )
    db.add_all([sup_m1, sup_m2])
    db.commit()

    # Build a test Excel workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Precios"
    ws.append(["Código", "Descripción", "Presentación", "Precio"])
    ws.append(["ACE-01", "Aceite Cañuelas 1.5L", "Botella", 1800])
    ws.append(["HAR-02", "Harina Pureza 1kg", "Paquete", 950])
    buf = io.BytesIO()
    wb.save(buf)
    excel_bytes = buf.getvalue()
    excel_b64 = base64.b64encode(excel_bytes).decode("utf-8")

    sent_messages = []

    async def mock_send(to_phone, text):
        sent_messages.append({"to_phone": to_phone, "text": text})
        return True

    with patch("app.services.whatsapp.send_whatsapp_message", side_effect=mock_send):
        res = client.post("/webhook", json={
            "phone": supplier_phone,
            "doc_bytes_b64": excel_b64,
            "doc_name": "Lista_Distribuidora_El_Progreso.xlsx"
        })
        assert res.status_code == 200
        assert res.json().get("action") == "supplier_catalog_loaded"
        assert res.json().get("count") >= 2

        import asyncio
        await asyncio.sleep(0.05)

        # Both merchants received catalog alert
        m1_msgs = [m for m in sent_messages if m["to_phone"] == m1_phone]
        m2_msgs = [m for m in sent_messages if m["to_phone"] == m2_phone]
        assert len(m1_msgs) >= 1
        assert "NUEVA LISTA DE PRECIOS DE TU PROVEEDOR" in m1_msgs[0]["text"]
        assert len(m2_msgs) >= 1
        assert "NUEVA LISTA DE PRECIOS DE TU PROVEEDOR" in m2_msgs[0]["text"]

        # Both merchants have products persisted in MerchantProduct
        m1_prods = db.query(MerchantProduct).filter(MerchantProduct.merchant_phone == m1_phone).all()
        m2_prods = db.query(MerchantProduct).filter(MerchantProduct.merchant_phone == m2_phone).all()
        assert len(m1_prods) >= 2
        assert len(m2_prods) >= 2


@pytest.mark.asyncio
async def test_client_deletion_on_the_fly(db):
    # 1. Test intent parsing for various formats
    p1 = parse_client_deletion_intent("Sofi, eliminar comercio Kiosco Alameda")
    assert p1["is_client_deletion"] is True
    assert p1["target_name"] == "Kiosco Alameda"

    p2 = parse_client_deletion_intent("sofia elimina a kiosco alameda")
    assert p2["is_client_deletion"] is True
    assert p2["target_name"].lower() == "kiosco alameda"

    p3 = parse_client_deletion_intent("dar de baja comercio al 3435551122")
    assert p3["is_client_deletion"] is True
    assert p3["phone"] == "3435551122"

    p4 = parse_client_deletion_intent("dar de baja mi comercio")
    assert p4["is_client_deletion"] is True
    assert p4["is_self"] is True

    # Negative guard: suppliers or employees are not clients
    assert parse_client_deletion_intent("eliminar proveedor Distribuidora Alem")["is_client_deletion"] is False
    assert parse_client_deletion_intent("eliminar empleado Lucas")["is_client_deletion"] is False

    # 2. Setup a client in DB with employee and draft orders
    with patch("app.services.whatsapp.send_whatsapp_message", new_callable=AsyncMock):
        # Onboard client
        h_onb, r_onb, a_onb = await process_boss_message(
            db,
            settings.WHATSAPP_ALERT_PHONE,
            "Sofi, dar de alta al comercio Kiosco Alameda de Juan al 3435551122 de Paraná"
        )
        assert h_onb is True
        assert a_onb == "client_onboarded"

        client_phone = "5493435551122"
        client_rec = db.query(Prospect).filter(Prospect.phone == client_phone).first()
        assert client_rec is not None
        assert client_rec.name == "Kiosco Alameda"

        # Add employee
        h_emp, r_emp, a_emp = await process_boss_message(
            db,
            client_phone,
            "Sofi, agregá a Lucas como empleado al 3435559999"
        )
        assert h_emp is True
        assert a_emp == "employee_added"

        # Add draft order
        draft = SupplierDraftOrder(
            merchant_phone=client_phone,
            supplier_key="distribuidora_test",
            supplier_name="Distribuidora Test",
            items=json.dumps([{"product_name": "Golosinas", "quantity": 10, "unit_price": 500}])
        )
        db.add(draft)
        db.commit()

        assert db.query(SupplierDraftOrder).filter(SupplierDraftOrder.merchant_phone == client_phone).count() == 1
        assert db.query(Prospect).filter(Prospect.parent_merchant_phone == client_phone).count() == 1

        # 3. Boss deletes the client by name
        h_del, r_del, a_del = await process_boss_message(
            db,
            settings.WHATSAPP_ALERT_PHONE,
            "Sofi, eliminar comercio Kiosco Alameda"
        )
        assert h_del is True
        assert a_del == "client_deleted"
        assert "COMERCIO DADO DE BAJA CON ÉXITO" in r_del
        assert "Kiosco Alameda" in r_del


        # Verify client and all associated test data are deleted
        assert db.query(Prospect).filter(Prospect.phone == client_phone).first() is None
        assert db.query(SupplierDraftOrder).filter(SupplierDraftOrder.merchant_phone == client_phone).count() == 0
        assert db.query(Prospect).filter(Prospect.parent_merchant_phone == client_phone).count() == 0


@pytest.mark.asyncio
async def test_supplier_phone_update_preserving_catalog(db):
    merchant_phone = "5493434111222"

    # 1. Test intent parsing for various formats
    p1 = await parse_supplier_phone_update_intent("Sofi, Carlos de Distribuidora Alem cambió de número al 3434552222")
    assert p1["is_supplier_phone_update"] is True
    assert p1["new_phone"] == "3434552222"
    assert "Alem" in p1["supplier_name"]

    p2 = await parse_supplier_phone_update_intent("actualizá el número de Distribuidora Alem al 3434552222")
    assert p2["is_supplier_phone_update"] is True
    assert p2["new_phone"] == "3434552222"

    p3 = await parse_supplier_phone_update_intent("Distribuidora Alem cambió de número, ahora es 3434552222")
    assert p3["is_supplier_phone_update"] is True
    assert p3["new_phone"] == "3434552222"

    # 2. Register supplier and populate cart
    with patch("app.services.whatsapp.send_whatsapp_message", new_callable=AsyncMock):
        # Register supplier
        await process_boss_message(
            db,
            merchant_phone,
            "Sofi, agendá a Carlos de Distribuidora Alem al 3434551111"
        )

        sup_rec = db.query(Prospect).filter(
            Prospect.merchant_phone == merchant_phone,
            Prospect.name == "Distribuidora Alem"
        ).first()
        assert sup_rec is not None
        assert sup_rec.phone == "5493434551111"

        # Add draft cart item for this supplier
        draft = SupplierDraftOrder(
            merchant_phone=merchant_phone,
            supplier_key="distribuidora_alem",
            supplier_name="Distribuidora Alem",
            items=json.dumps([{"product_name": "Cerveza Quilmes", "quantity": 5, "unit_price": 1200}])
        )
        db.add(draft)
        db.commit()

        # 3. Update supplier phone
        h_upd, r_upd, a_upd = await process_boss_message(
            db,
            merchant_phone,
            "Sofi, Carlos de Distribuidora Alem cambió de número al 3434552222"
        )
        assert h_upd is True
        assert a_upd == "supplier_phone_updated"
        assert "TELÉFONO DE PROVEEDOR ACTUALIZADO" in r_upd
        assert "Distribuidora Alem" in r_upd

        # 4. Verify supplier in DB has updated phone
        db.refresh(sup_rec)
        assert sup_rec.phone == "5493434552222"

        # 5. Verify draft basket is 100% intact
        saved_draft = db.query(SupplierDraftOrder).filter(
            SupplierDraftOrder.merchant_phone == merchant_phone,
            SupplierDraftOrder.supplier_key == "distribuidora_alem"
        ).first()
        assert saved_draft is not None
        items = json.loads(saved_draft.items)
        assert len(items) == 1
        assert items[0]["product_name"] == "Cerveza Quilmes"




