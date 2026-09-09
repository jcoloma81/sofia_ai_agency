import io
import pytest
import openpyxl
from app.services.catalog import CatalogService, clean_price, clean_stock, ProductItem

def test_clean_price():
    assert clean_price("$ 14.400") == 14400.0
    assert clean_price("14.400,50") == 14400.50
    assert clean_price("$14,400.00") == 14400.0
    assert clean_price("12500") == 12500.0
    assert clean_price(None) == 0.0

def test_clean_stock():
    assert clean_stock("SI") is True
    assert clean_stock("disponible") is True
    assert clean_stock("NO") is False
    assert clean_stock("sin stock") is False
    assert clean_stock("0") is False

def test_load_from_csv():
    csv_sample = """Codigo;Producto;Presentacion;Precio;Stock;Categoria
ART01;Aceite Cañuelas 1.5L;Caja x 6;14.400;SI;Almacén
ART02;Galletitas Oreo;Caja x 24;26400;NO;Golosinas
ART03;Harina Pureza 000;Fardo x 10;12500,00;SI;Almacén"""

    service = CatalogService()
    count = service.load_from_csv(csv_sample)
    assert count == 3
    assert len(service.products) == 3

    # Product 1
    p1 = service.find_product_exact_or_best("aceite")
    assert p1 is not None
    assert p1.price == 14400.0
    assert p1.in_stock is True
    assert p1.formatted_price() == "$14.400"

    # Product 2 (out of stock)
    p2 = service.find_product_exact_or_best("oreo")
    assert p2 is not None
    assert p2.in_stock is False

    # Search
    results = service.search_products("almacén")
    assert len(results) == 2

def test_load_from_excel_bytes():
    # Create workbook in-memory
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Código", "Artículo", "Bulto", "Precio Venta", "Stock"])
    ws.append(["A1", "Cerveza Quilmes 1L", "Cajón x 12", 24000, "SI"])
    ws.append(["A2", "Coca Cola 2.25L", "Pack x 6", 19200, "SI"])

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    service = CatalogService()
    count = service.load_from_excel_bytes(stream.read())
    assert count == 2
    
    p = service.find_product_exact_or_best("quilmes")
    assert p is not None
    assert p.price == 24000.0
    assert p.presentation == "Cajón x 12"

def test_google_sheet_url_normalization():
    edit_url = "https://docs.google.com/spreadsheets/d/1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms/edit#gid=0"
    export_url = CatalogService.normalize_google_sheet_url(edit_url)
    assert "export?format=csv&gid=0" in export_url
    assert "1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms" in export_url


def test_load_messy_excel_with_banners_and_sections():
    """Tests an Excel that has title logos, blank rows, section dividers and page break headers."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lista Caótica"

    # Row 1-4: Company banners, CUIT, date, blank
    ws.append(["DISTRIBUIDORA SAN JORGE S.A.", None, None, None])
    ws.append(["CUIT 30-71234567-8 - TEL: 11-4567-8900", None, None, None])
    ws.append(["LISTA OFICIAL SEPTIEMBRE 2026", None, None, None])
    ws.append([None, None, None, None])

    # Row 5: Real Table Header
    ws.append(["Cod. Art", "Descripción del Producto", "Formato Envase", "Precio Mayorista ($)", "Disp"])

    # Section 1 banner
    ws.append(["--- ALMACÉN Y SECOS ---", None, None, None, None])
    ws.append(["ALM01", "Arroz Gallo Oro 1kg", "Bolsa x 10", "$ 18.500,00", "SI"])
    ws.append(["ALM02", "Fideos Matarazzo 500g", "Caja x 20", "22400", "SI"])

    # Section 2 banner with different syntax
    ws.append(["RUBRO: BEBIDAS", None, None, None, None])
    ws.append(["BEB01", "Vino Malbec Norton 750ml", "Caja x 6", 36000.0, "SI"])

    # Repeated header from page break
    ws.append(["Cod. Art", "Descripción del Producto", "Formato Envase", "Precio Mayorista ($)", "Disp"])
    ws.append(["BEB02", "Fernet Branca 750ml", "Caja x 6", "$ 58.000.-", "NO"])

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    service = CatalogService()
    count = service.load_from_excel_bytes(stream.read(), filename="messy.xlsx")

    # Should have parsed exactly 4 products, skipping banners and repeated headers
    assert count == 4
    assert len(service.products) == 4

    p1 = service.find_product_exact_or_best("arroz")
    assert p1 is not None
    assert p1.price == 18500.0
    assert p1.presentation == "Bolsa x 10"
    assert p1.category == "ALMACÉN Y SECOS"
    assert p1.in_stock is True

    p3 = service.find_product_exact_or_best("norton")
    assert p3 is not None
    assert p3.price == 36000.0
    assert p3.category == "BEBIDAS"

    p4 = service.find_product_exact_or_best("branca")
    assert p4 is not None
    assert p4.price == 58000.0
    assert p4.in_stock is False


def test_multi_sheet_excel():
    """Tests an Excel workbook with multiple tabs/sheets for different product families."""
    wb = openpyxl.Workbook()
    
    # Sheet 1: Limpieza
    ws1 = wb.active
    ws1.title = "Limpieza"
    ws1.append(["Articulo", "Precio", "Pack"])
    ws1.append(["Lavandina Ayudin 1L", 8500, "Caja x 12"])
    ws1.append(["Detergente Magistral 500ml", 14200, "Caja x 12"])

    # Sheet 2: Golosinas
    ws2 = wb.create_sheet(title="Golosinas")
    ws2.append(["Articulo", "Precio", "Pack"])
    ws2.append(["Alfajor Jorgito Chocolate", 18000, "Caja x 24"])

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    service = CatalogService()
    count = service.load_from_excel_bytes(stream.read(), filename="multi_sheet.xlsx")

    assert count == 3
    lavandina = service.find_product_exact_or_best("lavandina")
    assert lavandina is not None
    assert lavandina.category == "Limpieza"

    alfajor = service.find_product_exact_or_best("jorgito")
    assert alfajor is not None
    assert alfajor.category == "Golosinas"


def test_supplier_excel_cross_update():
    """Tests intelligent cross-referencing and price updating from supplier Excel sheets."""
    service = CatalogService()

    # 1. Base catalog
    base_csv = """Codigo;Producto;Presentacion;Precio;Stock;Categoria
ART01;Aceite Cañuelas 1.5L;Caja x 6;14400;SI;Almacén
ART02;Harina Pureza 000 1kg;Fardo x 10;12500;SI;Almacén
ART03;Coca Cola 2.25L;Pack x 6;19200;SI;Bebidas"""
    service.load_from_csv(base_csv)
    assert len(service.products) == 3

    # 2. Supplier Excel with price hikes and natural naming differences
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["PROVEEDOR MOLINOS S.A. - LISTA DE PRECIOS", "", "", ""])
    ws.append(["", "", "", ""])
    ws.append(["Cód Prov", "Descripción de Fábrica", "Precio Nuevo", "Estado"])
    ws.append(["P-101", "Aceite Girasol Cañuelas 1500cc Botella", 16200, "Aumento"])
    ws.append(["P-102", "Harina Trigo Pureza 000 x 1000g", 13900, "Aumento"])
    ws.append(["P-103", "Gaseosa Coca-Cola Sabor Original 2250cc", 21500, "Aumento"])
    ws.append(["P-999", "Galletitas Chocolinas 250g", 9800, "Nuevo"])

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    res = service.update_from_supplier_excel(stream.read(), filename="aumentos_molinos.xlsx")
    assert res["status"] == "success"
    assert res["matched_count"] == 3
    assert res["new_count"] == 1

    # Verify updated prices in memory
    p_aceite = service.find_product_exact_or_best("aceite")
    assert p_aceite is not None
    assert p_aceite.price == 16200.0

    p_harina = service.find_product_exact_or_best("harina")
    assert p_harina is not None
    assert p_harina.price == 13900.0

    p_coca = service.find_product_exact_or_best("coca")
    assert p_coca is not None
    assert p_coca.price == 21500.0


