import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def create_styled_excel(file_path: str, sheet_title: str, headers: list, rows: list, header_bg="1B365D"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title

    # Header style
    h_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    h_fill = PatternFill(start_color=header_bg, end_color=header_bg, fill_type="solid")
    h_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Data style: explicit black text and white fill for 100% Android compatibility
    d_font = Font(name="Calibri", size=10, color="000000")
    d_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='D3D3D3'),
        right=Side(style='thin', color='D3D3D3'),
        top=Side(style='thin', color='D3D3D3'),
        bottom=Side(style='thin', color='D3D3D3')
    )

    # Write headers
    ws.append(headers)
    ws.row_dimensions[1].height = 26
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = h_font
        cell.fill = h_fill
        cell.alignment = h_align
        cell.border = thin_border

    # Write data
    for r_idx, row in enumerate(rows, start=2):
        ws.append(row)
        ws.row_dimensions[r_idx].height = 20
        for col_idx in range(1, len(row) + 1):
            cell = ws.cell(row=r_idx, column=col_idx)
            cell.font = d_font
            cell.fill = d_fill
            cell.border = thin_border
            # Format numbers/prices
            val = row[col_idx - 1]
            if isinstance(val, (int, float)):
                cell.number_format = '$ #,##0'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left" if col_idx > 1 else "center", vertical="center")

    # Column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    os.makedirs(os.path.dirname(file_path), exist_ok=True)
    wb.save(file_path)
    print(f"Generated: {file_path}")

def generate_all():
    base_dir = os.path.dirname(__file__)

    # --- FERRETERÍA ---
    ferr_dir = os.path.join(base_dir, "ejemplos excel ferreteria")
    
    # 1. Eléctrica Paraná
    create_styled_excel(
        file_path=os.path.join(ferr_dir, "1_proveedor_electrica_parana.xlsx"),
        sheet_title="Lista Oficial",
        headers=["CÓDIGO ARTÍCULO", "DESCRIPCIÓN DEL PRODUCTO", "MARCA", "UNIDAD", "PRECIO MAYORISTA"],
        rows=[
            ["ELE-009W", "Lámpara LED 9W Luz Cálida E27", "Osram", "Unidad", 950],
            ["ELE-DISC-115", "Disco corte fino metal 115mm", "Tyrolit", "Unidad", 1250],
            ["ELE-CINT-20", "Cinta aisladora negra 20 metros", "3M", "Rollo", 900],
            ["ELE-CAB-25", "Cable unipolar 2.5mm normalizado", "Prysmian", "Rollo x 100m", 38500],
            ["ELE-TERM-20", "Térmica bipolar 20A curva C", "Schneider", "Unidad", 8400]
        ],
        header_bg="0D47A1"
    )

    # 2. Distribuidora Nogoyá
    create_styled_excel(
        file_path=os.path.join(ferr_dir, "2_proveedor_distribuidora_nogoya.xlsx"),
        sheet_title="Precios Vigentes",
        headers=["SKU", "PRODUCTO", "RUBRO", "PRESENTACIÓN", "PRECIO LISTA"],
        rows=[
            ["NOG-8812", "Foco LED 9W Luz Fría E27", "Electricidad", "Unidad", 1150],
            ["NOG-DISC-01", "Disco de corte amoladora 115mm x 1mm", "Abrasivos", "Unidad", 1400],
            ["NOG-TORN-10", "Tornillo autoperforante 10x3/4 cabeza tanque", "Tornillería", "Caja x 1000", 7900],
            ["NOG-PIN-08", "Pinza universal 8 pulgadas aislada", "Herramientas", "Unidad", 12500],
            ["NOG-ALI-06", "Alicate corte diagonal 6 pulgadas", "Herramientas", "Unidad", 11000],
            ["NOG-MART-500", "Martillo galponero mango fibra 500g", "Herramientas", "Unidad", 14500]
        ],
        header_bg="1B5E20"
    )

    # 3. Mayorista Central
    create_styled_excel(
        file_path=os.path.join(ferr_dir, "3_proveedor_mayorista_central.xlsx"),
        sheet_title="Tarifa Mayorista",
        headers=["COD_ITEM", "DETALLE", "FAMILIA", "EMPAQUE", "P.UNITARIO"],
        rows=[
            ["CEN-991", "Lámpara LED 12W Luz Fría E27", "Iluminación", "Unidad", 1450],
            ["CEN-AMO-850", "Amoladora angular 115mm 850W", "Máquinas", "Unidad", 68000],
            ["CEN-TAL-650", "Taladro percutor 13mm 650W", "Máquinas", "Unidad", 74000]
        ],
        header_bg="B71C1C"
    )

    # --- ALIMENTOS / DESPENSA ---
    alim_dir = os.path.join(base_dir, "ejemplos excel alimentos")

    # 1. Molinos Cañuelas
    create_styled_excel(
        file_path=os.path.join(alim_dir, "1_proveedor_molinos_canuelas.xlsx"),
        sheet_title="Lista Fábrica",
        headers=["CÓDIGO", "PRODUCTO", "PRESENTACIÓN", "PRECIO X UNIDAD"],
        rows=[
            ["MC-101", "Harina 000 Cañuelas 1kg", "Fardo x 10", 1150],
            ["MC-205", "Aceite Cañuelas 1.5L", "Caja x 6", 2300],
            ["MC-304", "Galletitas 9 de Oro agridulces 200g", "Caja x 24", 800]
        ],
        header_bg="E65100"
    )

    # 2. Distribuidora San Martín
    create_styled_excel(
        file_path=os.path.join(alim_dir, "2_proveedor_distribuidora_san_martin.xlsx"),
        sheet_title="Catálogo Distribución",
        headers=["ART", "DESCRIPCION_ARTICULO", "RUBRO", "PACK", "PRECIO_MAYORISTA"],
        rows=[
            ["SM-779", "Harina Pureza 000 1kg", "Almacén", "Fardo x 10", 1250],
            ["SM-900", "Aceite de Girasol Natura 900ml", "Almacén", "Caja x 12", 1850],
            ["SM-520", "Puré de Tomate La Campagnola 520g", "Conservas", "Caja x 12", 980],
            ["SM-475", "Mayonesa Natura doy pack 500g", "Aderezos", "Caja x 12", 1450],
            ["SM-101", "Azúcar Ledesma Clásica 1kg", "Endulzantes", "Fardo x 10", 1050]
        ],
        header_bg="37474F"
    )

    # 3. Mayorista Litoral
    create_styled_excel(
        file_path=os.path.join(alim_dir, "3_proveedor_mayorista_litoral.xlsx"),
        sheet_title="Ofertas Mayoristas",
        headers=["SKU", "PRODUCTO", "CATEGORÍA", "FORMATO", "PRECIO_UNITARIO"],
        rows=[
            ["ML-045", "Harina 000 Morixe 1kg", "Harinas", "Fardo x 10", 1320],
            ["ML-112", "Aceite Cocinero Girasol 900ml", "Aceites", "Caja x 12", 1790],
            ["ML-201", "Arroz Lucchetti Largo Fino 1kg", "Granos", "Fardo x 10", 1600]
        ],
        header_bg="4A148C"
    )

if __name__ == "__main__":
    generate_all()
